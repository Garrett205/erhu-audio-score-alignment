import os
import re
import math
import itertools
import numpy as np
import pandas as pd
from openpyxl.styles import PatternFill

from common_utils import (
    cents_error,
    classify_duration_ratio_level,
    classify_pitch_level,
    classify_rhythm_level,
    find_first_with_suffix,
    freq_to_midi,
    safe_float,
    safe_numeric,
)
from alignment_contract import (
    CURRENT_ONSET_INDEX_COL,
    ONSET_ID_COL,
    ORIGINAL_ONSET_INDEX_COL,
    choose_onset_file,
    ensure_onset_identity,
    finalize_onset_dataframe,
    onset_index_by_id,
    split_normal_and_weak_onsets,
)

# =========================
# 全局配置
# =========================
RUN_TAG = "DP2"
METHOD_TAG_CN = "score主导"

# ===== 文件名偏好 =====
DP1_ALIGN_KEYWORD = "_DP1_对齐结果_"
DP1_METRIC_KEYWORD = "_DP1_指标统计_"

# ===== 列名（score） =====
SCORE_TIME_COL   = "累计时间(s)"
SCORE_F0_COL     = "音高频率(Hz)"
SCORE_NAME_COL   = "音名"
SCORE_DUR_COL    = "持续时间(s)"
SCORE_BPM_COL    = "BPM"
SCORE_MISC_COL   = "连音线和其他信息"
SCORE_LEGATO_COL = "连奏信息"
SCORE_GRACE_COL  = "装饰音判定"

# ===== 列名（onset） =====
ONSET_F0_COL = "note_f0(Hz)"
ONSET_T_COL  = "note_time(s)"

# ===== onset 去重 =====
ENABLE_ONSET_DEDUP = True
GRACE_GAP_RATIO_TOL = 0.6

# ===== DP2 模块修复参数 =====
MODULE_CAND_CENTS = 100.0
MODULE_SKIP_ONSET_PENALTY = 0.8
MODULE_EARLY_WEIGHT = 3.5
MODULE_PITCH_WEIGHT = 0.9
MODULE_PROGRESS_WEIGHT = 0.5

# ===== 单点 F 修复参数 =====
SINGLE_PITCH_SIGMA = 40.0
SINGLE_BAND_RATIO = 0.72
SINGLE_TIME_RATIO_TOL = 2.5
SINGLE_ACCEPT_MIN_CENT_IMPROVE = 12.0
SINGLE_ACCEPT_TIME_MS_TOL = 80.0

# ===== 节奏 F 三音块重排参数 =====
TRIPLET_PITCH_CENTS_MAX = 120.0
TRIPLET_DUR_WEIGHT = 1.0
TRIPLET_PITCH_WEIGHT = 0.20
TRIPLET_TARGET_WEIGHT = 0.06
TRIPLET_ACCEPT_IMPROVE = 0.12
TRIPLET_MAX_CANDS_PER_NOTE = 8

DP2_ACTIVE_REPAIR_STAGES = ("坏区边界恢复", "坏区主体模块修复")
DP2_DEFERRED_TO_DP3 = ("单点F修复", "节奏F三音重排")


def gaussian_pitch_score(abs_cents, sigma=SINGLE_PITCH_SIGMA):
    return float(np.exp(- (float(abs_cents) / sigma) ** 2))


def nearest_success_left(df, idx):
    for k in range(idx - 1, -1, -1):
        if str(df.at[k, "对齐状态"]) == "对齐成功":
            return k
    return None


def nearest_success_right(df, idx):
    for k in range(idx + 1, len(df)):
        if str(df.at[k, "对齐状态"]) == "对齐成功":
            return k
    return None


def fold_freq_to_ref(freq, ref_freq):
    """
    去掉八度调整：直接返回原频率。
    保留函数名仅为了兼容原有调用结构。
    """
    freq = float(freq)
    return freq, abs(cents_error(ref_freq, freq))


def same_pitch_group(freq1, freq2):
    m1 = int(round(freq_to_midi([freq1])[0]))
    m2 = int(round(freq_to_midi([freq2])[0]))
    return m1 == m2


def parse_bad_region_tag(tag):
    if not isinstance(tag, str) or "坏区" not in tag:
        return None
    m = re.search(r"原可疑:(\d+)-(\d+);\s*删除:(\d+)-(\d+)", tag)
    if not m:
        return None
    return {
        "raw_start": int(m.group(1)),
        "raw_end": int(m.group(2)),
        "drop_start": int(m.group(3)),
        "drop_end": int(m.group(4)),
    }


def extract_bad_regions_from_table(df):
    regions = []
    seen = set()
    if "可疑坏区间标注" not in df.columns:
        return regions
    for _, tag in enumerate(df["可疑坏区间标注"].fillna("")):
        info = parse_bad_region_tag(tag)
        if info is None:
            continue
        key = (info["raw_start"], info["raw_end"], info["drop_start"], info["drop_end"])
        if key in seen:
            continue
        seen.add(key)
        info["tag"] = tag
        regions.append(info)
    regions.sort(key=lambda x: x["raw_start"])
    return regions


def find_score_file(base_dir):
    return find_first_with_suffix(base_dir, "_score_final.csv", "*_score_final.csv")


def find_onset_file(base_dir):
    return choose_onset_file(base_dir).name


def find_dp1_align_file(base_dir):
    cand = [f for f in os.listdir(base_dir) if DP1_ALIGN_KEYWORD in f and f.endswith(".csv")]
    if len(cand) == 0:
        raise RuntimeError("未找到 DP1 对齐结果 CSV（文件名需包含 _DP1_对齐结果_）")
    if len(cand) > 1:
        print(f"⚠️ 检测到多个 DP1 结果文件，默认使用：{cand[0]}")
    return cand[0]


def load_data(base_dir):
    score_file = find_score_file(base_dir)
    onset_file = find_onset_file(base_dir)
    dp1_file = find_dp1_align_file(base_dir)

    df_sc = pd.read_csv(os.path.join(base_dir, score_file))
    df_on = ensure_onset_identity(pd.read_csv(os.path.join(base_dir, onset_file)))
    df_on, _weak_onsets = split_normal_and_weak_onsets(df_on)
    df_dp1 = pd.read_csv(os.path.join(base_dir, dp1_file))

    for c in [SCORE_TIME_COL, SCORE_F0_COL, SCORE_DUR_COL]:
        if c not in df_sc.columns:
            raise RuntimeError(f"score 文件缺少列: {c}")
    for c in [ONSET_T_COL, ONSET_F0_COL]:
        if c not in df_on.columns:
            raise RuntimeError(f"onset 文件缺少列: {c}")

    df_sc[SCORE_TIME_COL] = safe_numeric(df_sc[SCORE_TIME_COL])
    df_sc[SCORE_F0_COL] = safe_numeric(df_sc[SCORE_F0_COL])
    df_sc[SCORE_DUR_COL] = safe_numeric(df_sc[SCORE_DUR_COL])
    df_on[ONSET_T_COL] = safe_numeric(df_on[ONSET_T_COL])
    df_on[ONSET_F0_COL] = safe_numeric(df_on[ONSET_F0_COL])

    df_sc = df_sc.dropna(subset=[SCORE_TIME_COL, SCORE_F0_COL]).reset_index(drop=True)
    df_on = df_on.dropna(subset=[ONSET_T_COL, ONSET_F0_COL]).sort_values(ONSET_T_COL).reset_index(drop=True)

    if ENABLE_ONSET_DEDUP:
        df_on["_t_round"] = df_on[ONSET_T_COL].round(4)
        before = len(df_on)
        df_on = df_on.drop_duplicates(subset=["_t_round"], keep="first").drop(columns=["_t_round"]).reset_index(drop=True)
        after = len(df_on)
        print(f"🧹 onset 时间去重: {before} -> {after}")

    df_on = finalize_onset_dataframe(df_on)

    prefix = os.path.splitext(score_file)[0].replace("_score_final", "")
    return df_sc, df_on, df_dp1, prefix, score_file, onset_file, dp1_file


def ensure_columns(df):
    default_cols = {
        "DP2修复来源": "",
        "DP2修复说明": "",
        "DP2候选onset_idx": np.nan,
        ONSET_ID_COL: "",
        ORIGINAL_ONSET_INDEX_COL: np.nan,
        CURRENT_ONSET_INDEX_COL: np.nan,
        "红点频率_raw(Hz)": np.nan,
        "红点频率_used(Hz)": np.nan,
        "低音八度修正": "否",
        "高八度等效匹配": "否",
        "对齐状态": "",
        "后删标记": "",
        "可疑点标记": "",
        "坏点标记": "",
        "可疑坏区间标注": "",
    }
    for c, v in default_cols.items():
        if c not in df.columns:
            df[c] = v
    return df


def _onset_time_array(df_on):
    return safe_numeric(df_on[ONSET_T_COL]).to_numpy(dtype=float)


def _legacy_exact_onset_idx(t, df_on):
    """Historical tables may omit onset_id; accept only one exact time match."""
    value = safe_float(t)
    if not np.isfinite(value):
        return np.nan
    times = _onset_time_array(df_on)
    matches = np.where(np.isclose(times, value, atol=1e-9, rtol=0.0))[0]
    return int(matches[0]) if len(matches) == 1 else np.nan


def build_onset_index_mapping(df_align, df_on, onset_times=None):
    if onset_times is None:
        onset_times = _onset_time_array(df_on)
    mapping = []
    for _, row in df_align.iterrows():
        if str(row.get("对齐状态", "")) != "对齐成功":
            mapping.append(np.nan)
            continue
        stable_id = str(row.get(ONSET_ID_COL, "")).strip()
        if stable_id:
            mapped = onset_index_by_id(df_on, stable_id)
            if mapped is None:
                raise ValueError(f"DP 输入引用了不存在的 onset_id: {stable_id}")
        else:
            # Compatibility only for historical DP1 tables: exact and unique,
            # never a nearest-time guess.
            mapped = _legacy_exact_onset_idx(row.get("红点时间(s)"), df_on)
        mapping.append(mapped)
    df_align["DP2候选onset_idx"] = mapping
    for row_index, mapped in enumerate(mapping):
        if np.isfinite(safe_float(mapped)):
            on = df_on.loc[int(mapped)]
            df_align.at[row_index, ONSET_ID_COL] = on[ONSET_ID_COL]
            df_align.at[row_index, ORIGINAL_ONSET_INDEX_COL] = on[ORIGINAL_ONSET_INDEX_COL]
            df_align.at[row_index, CURRENT_ONSET_INDEX_COL] = int(mapped)
    return df_align


def overwrite_match_at(df, idx, on_idx, df_on, note="", source=""):
    on = df_on.loc[int(on_idx)]
    red_t = float(on[ONSET_T_COL])
    red_f_raw = float(on[ONSET_F0_COL])

    df.at[idx, "红点时间(s)"] = red_t
    df.at[idx, "红点频率_raw(Hz)"] = red_f_raw
    df.at[idx, "红点频率_used(Hz)"] = red_f_raw
    df.at[idx, "DP2候选onset_idx"] = float(on_idx)
    df.at[idx, ONSET_ID_COL] = on[ONSET_ID_COL]
    df.at[idx, ORIGINAL_ONSET_INDEX_COL] = on[ORIGINAL_ONSET_INDEX_COL]
    df.at[idx, CURRENT_ONSET_INDEX_COL] = int(on_idx)
    df.at[idx, "低音八度修正"] = "否"
    df.at[idx, "高八度等效匹配"] = "否"
    df.at[idx, "对齐状态"] = "对齐成功"
    df.at[idx, "后删标记"] = ""
    df.at[idx, "坏点标记"] = ""
    if source:
        df.at[idx, "DP2修复来源"] = source
    if note:
        old = str(df.at[idx, "DP2修复说明"]).strip()
        df.at[idx, "DP2修复说明"] = (old + " | " + note).strip(" |")
    return df


def restore_bad_region_boundaries(df_align, df_on, onset_times=None):
    if onset_times is None:
        onset_times = _onset_time_array(df_on)
    restored = 0
    for reg in extract_bad_regions_from_table(df_align):
        for idx in [reg["drop_start"], reg["drop_end"]]:
            if idx < 0 or idx >= len(df_align):
                continue
            if reg["raw_start"] <= idx <= reg["raw_end"]:
                continue

            red_t = safe_float(df_align.at[idx, "红点时间(s)"])
            raw_f = safe_float(df_align.at[idx, "红点频率_raw(Hz)"])

            if not np.isfinite(raw_f) and np.isfinite(red_t):
                on_idx = _legacy_exact_onset_idx(red_t, df_on)
                if np.isfinite(on_idx):
                    raw_f = float(df_on.at[int(on_idx), ONSET_F0_COL])
                    df_align.at[idx, "红点频率_raw(Hz)"] = raw_f
                    df_align.at[idx, "DP2候选onset_idx"] = on_idx

            if np.isfinite(raw_f):
                df_align.at[idx, "红点频率_used(Hz)"] = raw_f
                df_align.at[idx, "低音八度修正"] = "否"
                df_align.at[idx, "高八度等效匹配"] = "否"

            if np.isfinite(red_t):
                df_align.at[idx, "对齐状态"] = "对齐成功"
                df_align.at[idx, "后删标记"] = ""
                df_align.at[idx, "坏点标记"] = ""
                old = str(df_align.at[idx, "DP2修复说明"]).strip()
                extra = "坏区边界恢复为对齐成功"
                df_align.at[idx, "DP2修复说明"] = (old + " | " + extra).strip(" |")
                if str(df_align.at[idx, "DP2修复来源"]).strip() == "":
                    df_align.at[idx, "DP2修复来源"] = "坏区边界恢复"
                if not np.isfinite(safe_float(df_align.at[idx, "DP2候选onset_idx"])):
                    df_align.at[idx, "DP2候选onset_idx"] = _legacy_exact_onset_idx(red_t, df_on)
                restored += 1
    return df_align, restored


def recalc_table_metrics(df_align):
    df = df_align.copy()
    score_shift = 0.0
    if "score_shift(s)" in df.columns and len(df) > 0:
        v = safe_numeric(df["score_shift(s)"]).dropna()
        if len(v) > 0:
            score_shift = float(v.iloc[0])

    if "谱面时间_原始(s)" not in df.columns and "累计时长(s)" in df.columns:
        df["谱面时间_原始(s)"] = safe_numeric(df["累计时长(s)"])
    df["谱面时间_平移后(s)"] = safe_numeric(df["谱面时间_原始(s)"]) + score_shift
    df["score_shift(s)"] = score_shift

    for i in range(len(df)):
        if str(df.at[i, "对齐状态"]) != "对齐成功":
            continue
        sc_t = safe_float(df.at[i, "谱面时间_平移后(s)"])
        red_t = safe_float(df.at[i, "红点时间(s)"])
        score_f = safe_float(df.at[i, "谱面频率(Hz)"])
        raw_f = safe_float(df.at[i, "红点频率_raw(Hz)"])
        used_f = safe_float(df.at[i, "红点频率_used(Hz)"])

        if np.isfinite(score_f) and np.isfinite(raw_f) and (not np.isfinite(used_f) or used_f <= 0):
            used_f = raw_f
            df.at[i, "红点频率_used(Hz)"] = used_f
            df.at[i, "低音八度修正"] = "否"
            df.at[i, "高八度等效匹配"] = "否"

        if np.isfinite(sc_t) and np.isfinite(red_t):
            time_err_ms = (red_t - sc_t) * 1000.0
            df.at[i, "时间误差(ms)"] = time_err_ms
            df.at[i, "节奏等级"] = classify_rhythm_level(time_err_ms)
        if np.isfinite(score_f) and np.isfinite(used_f) and score_f > 0 and used_f > 0:
            ce = cents_error(score_f, used_f)
            df.at[i, "音准误差(cents)"] = ce
            df.at[i, "音准等级"] = classify_pitch_level(ce)

    actual_gaps = []
    score_gaps = []
    ratios = []
    ratio_pcts = []
    ratio_levels = []

    # 保留用户修正版：时值统计沿用平移后的谱面时间，与 DP1 输出口径一致。
    score_times = safe_numeric(df["谱面时间_平移后(s)"]).to_numpy(dtype=float)
    onset_times = safe_numeric(df["红点时间(s)"]).to_numpy(dtype=float)
    score_durs = safe_numeric(df["持续时间(s)"]).to_numpy(dtype=float) if "持续时间(s)" in df.columns else np.full(len(df), np.nan)
    bpm_vals = safe_numeric(df["BPM"]).to_numpy(dtype=float) if "BPM" in df.columns else np.full(len(df), np.nan)
    grace_flags = df["装饰音判定"].fillna("").astype(str).to_list() if "装饰音判定" in df.columns else [""] * len(df)

    last_black_time = np.nan
    if len(df) > 0 and np.isfinite(score_times[-1]) and np.isfinite(score_durs[-1]):
        last_black_time = score_times[-1] + score_durs[-1]

    statuses = df["对齐状态"].astype(str).to_numpy()
    next_success_idx = np.full(len(df), -1, dtype=int)
    next_success = -1
    for i in range(len(df) - 1, -1, -1):
        next_success_idx[i] = next_success
        if statuses[i] == "对齐成功":
            next_success = i

    for i in range(len(df)):
        if statuses[i] != "对齐成功":
            actual_gaps.append(np.nan)
            score_gaps.append(np.nan)
            ratios.append(np.nan)
            ratio_pcts.append(np.nan)
            ratio_levels.append("F")
            continue

        next_success = next_success_idx[i]
        if next_success >= 0:
            actual_gap = onset_times[next_success] - onset_times[i]
            score_gap = score_times[next_success] - score_times[i]
        else:
            actual_gap = last_black_time - onset_times[i] if np.isfinite(last_black_time) and np.isfinite(onset_times[i]) else np.nan
            is_grace = grace_flags[i].strip() == "装饰音"
            if is_grace:
                bpm_val = bpm_vals[i]
                score_gap = 6.0 / bpm_val if np.isfinite(bpm_val) and bpm_val > 0 else np.nan
            else:
                score_gap = score_durs[i] if np.isfinite(score_durs[i]) else np.nan

        ratio = actual_gap / score_gap if np.isfinite(actual_gap) and np.isfinite(score_gap) and abs(score_gap) > 1e-12 else np.nan
        actual_gaps.append(actual_gap)
        score_gaps.append(score_gap)
        ratios.append(ratio)
        ratio_pcts.append((ratio - 1.0) * 100.0 if np.isfinite(ratio) else np.nan)
        ratio_levels.append(classify_duration_ratio_level(ratio, grace_flags[i].strip() == "装饰音"))

    df["实际时值(s)"] = actual_gaps
    df["谱面时值(s)"] = score_gaps
    df["时值比"] = ratios
    df["时值偏差(%)"] = ratio_pcts
    df["时值等级"] = ratio_levels
    return df


def build_score_modules(df_sc, raw_l, raw_r):
    modules = []
    cur = None
    for idx in range(raw_l, raw_r + 1):
        f = float(df_sc.at[idx, SCORE_F0_COL])
        dur = float(df_sc.at[idx, SCORE_DUR_COL]) if pd.notna(df_sc.at[idx, SCORE_DUR_COL]) else 0.0
        if cur is None:
            cur = {"start": idx, "end": idx, "ref_freq": f, "indices": [idx], "durations": [dur]}
            continue
        if same_pitch_group(cur["ref_freq"], f):
            cur["end"] = idx
            cur["indices"].append(idx)
            cur["durations"].append(dur)
        else:
            modules.append(cur)
            cur = {"start": idx, "end": idx, "ref_freq": f, "indices": [idx], "durations": [dur]}
    if cur is not None:
        modules.append(cur)
    return modules


def local_module_dp(modules, cand_indices, df_on):
    m, k = len(modules), len(cand_indices)
    if m == 0 or k == 0:
        return None
    cand_times = safe_numeric(df_on.loc[cand_indices, ONSET_T_COL]).to_numpy(dtype=float)
    cand_freqs = safe_numeric(df_on.loc[cand_indices, ONSET_F0_COL]).to_numpy(dtype=float)
    INF = 1e18
    dp = np.full((m + 1, k + 1), INF, dtype=float)
    prev = np.empty((m + 1, k + 1), dtype=object)
    dp[0, :] = np.arange(k + 1) * MODULE_SKIP_ONSET_PENALTY
    prev[0, 0] = None
    for j in range(1, k + 1):
        prev[0, j] = ("skip", 0, j - 1)

    for i in range(1, m + 1):
        mod = modules[i - 1]
        ref_f = float(mod["ref_freq"])
        for j in range(1, k + 1):
            best = dp[i, j - 1] + MODULE_SKIP_ONSET_PENALTY
            best_prev = ("skip", i, j - 1)
            on_f = cand_freqs[j - 1]
            on_f_used = on_f
            ce_abs = abs(cents_error(ref_f, on_f_used))
            if ce_abs <= MODULE_CAND_CENTS:
                pitch_cost = ce_abs / 100.0
                early_cost = (j - 1)
                progress_cost = abs(j / max(k, 1) - i / max(m, 1))
                cand = dp[i - 1, j - 1] + MODULE_PITCH_WEIGHT * pitch_cost + MODULE_EARLY_WEIGHT * early_cost / max(k, 1) + MODULE_PROGRESS_WEIGHT * progress_cost
                if cand < best:
                    best = cand
                    best_prev = ("match", i - 1, j - 1, float(on_f_used), float(ce_abs))
            dp[i, j] = best
            prev[i, j] = best_prev

    j_best = int(np.argmin(dp[m, :]))
    assign = [None] * m
    i, j = m, j_best
    while i > 0 and j >= 0:
        p = prev[i, j]
        if p is None:
            break
        if p[0] == "skip":
            _, i2, j2 = p
            i, j = i2, j2
        else:
            _, mi, cj, on_f_used, ce_abs = p
            assign[mi] = {
                "cand_local_idx": cj,
                "onset_idx": int(cand_indices[cj]),
                "onset_time": float(cand_times[cj]),
                "onset_f_raw": float(cand_freqs[cj]),
                "onset_f_used": float(on_f_used),
                "abs_cents": float(ce_abs),
            }
            i, j = mi, cj
    if any(x is None for x in assign):
        return None
    return assign


def split_module_by_duration(module, start_time, end_time, cand_indices, df_on):
    indices = module["indices"]
    durs = [float(max(d, 1e-9)) for d in module["durations"]]
    total_dur = float(sum(durs)) if len(durs) > 0 else 1.0
    if not np.isfinite(start_time) or not np.isfinite(end_time) or end_time <= start_time:
        end_time = start_time + total_dur
    targets = [start_time]
    cum = 0.0
    for d in durs[:-1]:
        cum += d
        targets.append(start_time + (end_time - start_time) * (cum / total_dur))

    cand_times = safe_numeric(df_on.loc[cand_indices, ONSET_T_COL]).to_numpy(dtype=float) if len(cand_indices) > 0 else np.array([])
    used = []
    last_pos = -1
    for p, target in enumerate(targets):
        if len(cand_indices) > 0:
            search_from = max(last_pos + 1, 0)
            if search_from >= len(cand_indices):
                used.append({"score_idx": indices[p], "onset_idx": np.nan, "time": float(target), "f_raw": np.nan})
                continue
            best_local = search_from + int(np.argmin(np.abs(cand_times[search_from:] - target)))
            last_pos = best_local
            used.append({
                "score_idx": indices[p],
                "onset_idx": int(cand_indices[best_local]),
                "time": float(df_on.at[int(cand_indices[best_local]), ONSET_T_COL]),
                "f_raw": float(df_on.at[int(cand_indices[best_local]), ONSET_F0_COL]),
            })
        else:
            used.append({"score_idx": indices[p], "onset_idx": np.nan, "time": float(target), "f_raw": np.nan})
    return used


def mark_unmatched_at(df, idx, reason):
    for col in ["红点时间(s)", "红点频率_raw(Hz)", "红点频率_used(Hz)", "DP2候选onset_idx",
                ORIGINAL_ONSET_INDEX_COL, CURRENT_ONSET_INDEX_COL, "时间误差(ms)", "音准误差(cents)"]:
        if col in df.columns:
            df.at[idx, col] = np.nan
    df.at[idx, ONSET_ID_COL] = ""
    df.at[idx, "对齐状态"] = "unmatched"
    df.at[idx, "DP2修复来源"] = "候选不足"
    df.at[idx, "DP2修复说明"] = reason
    return df


def enforce_unique_onset_assignments(df):
    """Keep the first score assignment and mark all later reuse as conflicts."""
    seen = set()
    for idx in range(len(df)):
        if str(df.at[idx, "对齐状态"]) != "对齐成功":
            continue
        onset_id = str(df.at[idx, ONSET_ID_COL]).strip()
        if not onset_id:
            continue
        if onset_id in seen:
            mark_unmatched_at(df, idx, f"onset_id 重复分配冲突: {onset_id}")
            df.at[idx, "对齐状态"] = "conflict"
        else:
            seen.add(onset_id)
    return df


def repair_bad_region_bodies(df_align, df_sc, df_on):
    repaired_rows = 0
    for reg in extract_bad_regions_from_table(df_align):
        raw_l, raw_r = reg["raw_start"], reg["raw_end"]
        left_anchor = nearest_success_left(df_align, raw_l)
        right_anchor = nearest_success_right(df_align, raw_r)
        if left_anchor is None or right_anchor is None:
            continue
        left_on_idx = safe_float(df_align.at[left_anchor, "DP2候选onset_idx"])
        right_on_idx = safe_float(df_align.at[right_anchor, "DP2候选onset_idx"])
        if not np.isfinite(left_on_idx) or not np.isfinite(right_on_idx):
            continue
        left_on_idx, right_on_idx = int(left_on_idx), int(right_on_idx)
        if right_on_idx - left_on_idx <= 1:
            continue

        modules = build_score_modules(df_sc, raw_l, raw_r)
        cand_indices = list(range(left_on_idx + 1, right_on_idx))
        assign = local_module_dp(modules, cand_indices, df_on)
        if assign is None:
            continue

        expanded = []
        for mi, mod in enumerate(modules):
            t0 = assign[mi]["onset_time"]
            t1 = assign[mi + 1]["onset_time"] if mi < len(modules) - 1 else float(df_align.at[right_anchor, "红点时间(s)"])
            local_cands = [c for c in cand_indices if t0 - 1e-9 <= float(df_on.at[c, ONSET_T_COL]) < t1 + 1e-9]
            expanded.extend(split_module_by_duration(mod, t0, t1, local_cands, df_on))

        for item in expanded:
            idx = item["score_idx"]
            if np.isfinite(item["onset_idx"]):
                overwrite_match_at(df_align, idx, int(item["onset_idx"]), df_on, note=f"坏区主体模块修复[{raw_l}-{raw_r}]", source="坏区主体模块修复")
                old_tag = str(df_align.at[idx, "可疑坏区间标注"]).strip()
                if old_tag and "主体已DP2修复" not in old_tag:
                    df_align.at[idx, "可疑坏区间标注"] = old_tag + " | 主体已DP2修复"
                repaired_rows += 1
            else:
                mark_unmatched_at(df_align, idx, f"坏区主体候选不足[{raw_l}-{raw_r}]")
    return enforce_unique_onset_assignments(df_align), repaired_rows


def macro_time_ok(df_align, left_idx, mid_time, right_idx, cur_idx):
    t_l = float(df_align.at[left_idx, "红点时间(s)"])
    t_r = float(df_align.at[right_idx, "红点时间(s)"])
    if not (np.isfinite(t_l) and np.isfinite(t_r) and t_r > t_l and mid_time > t_l and mid_time < t_r):
        return False
    s_l = float(df_align.at[left_idx, "谱面时间_原始(s)"])
    s_c = float(df_align.at[cur_idx, "谱面时间_原始(s)"])
    s_r = float(df_align.at[right_idx, "谱面时间_原始(s)"])
    d1 = max(s_c - s_l, 1e-9)
    d2 = max(s_r - s_c, 1e-9)
    r1 = (mid_time - t_l) / d1
    r2 = (t_r - mid_time) / d2
    return (1 / SINGLE_TIME_RATIO_TOL <= r1 <= SINGLE_TIME_RATIO_TOL) and (1 / SINGLE_TIME_RATIO_TOL <= r2 <= SINGLE_TIME_RATIO_TOL)


def repair_single_pitch_f(df_align, df_on):
    repaired = 0
    for i in range(len(df_align)):
        if str(df_align.at[i, "坏点标记"]) == "坏点":
            continue
        if str(df_align.at[i, "音准等级"]) != "F":
            continue
        left_idx = nearest_success_left(df_align, i)
        right_idx = nearest_success_right(df_align, i)
        if left_idx is None or right_idx is None:
            continue
        l_on = safe_float(df_align.at[left_idx, "DP2候选onset_idx"])
        r_on = safe_float(df_align.at[right_idx, "DP2候选onset_idx"])
        if not np.isfinite(l_on) or not np.isfinite(r_on):
            continue
        l_on, r_on = int(l_on), int(r_on)
        if r_on - l_on <= 1:
            continue

        score_f = float(df_align.at[i, "谱面频率(Hz)"])
        old_used = safe_float(df_align.at[i, "红点频率_used(Hz)"])
        old_time = safe_float(df_align.at[i, "红点时间(s)"])
        old_abs = abs(cents_error(score_f, old_used)) if np.isfinite(old_used) and old_used > 0 else 1e9

        cand_rows = []
        for on_idx in range(l_on + 1, r_on):
            on_t = float(df_on.at[on_idx, ONSET_T_COL])
            on_f = float(df_on.at[on_idx, ONSET_F0_COL])
            used_f = on_f
            abs_c = abs(cents_error(score_f, used_f))
            pscore = gaussian_pitch_score(abs_c)
            cand_rows.append((on_idx, on_t, on_f, used_f, abs_c, pscore))
        if not cand_rows:
            continue

        best_ps = max(x[5] for x in cand_rows)
        reasonable = [x for x in cand_rows if x[5] >= SINGLE_BAND_RATIO * best_ps]
        reasonable = [x for x in reasonable if macro_time_ok(df_align, left_idx, x[1], right_idx, i)]
        if not reasonable:
            continue
        reasonable.sort(key=lambda x: (x[1], x[4]))
        best = reasonable[0]
        new_abs = best[4]
        new_time = best[1]

        improve_pitch = old_abs - new_abs
        better_time = (new_time <= old_time + SINGLE_ACCEPT_TIME_MS_TOL / 1000.0) if np.isfinite(old_time) else True
        if (improve_pitch >= SINGLE_ACCEPT_MIN_CENT_IMPROVE and better_time) or (new_abs <= 100.0 < old_abs):
            overwrite_match_at(df_align, i, int(best[0]), df_on, note=f"单点F修复: {old_abs:.1f}c -> {new_abs:.1f}c", source="单点F修复")
            repaired += 1
    return df_align, repaired


def _triplet_block_valid(df_align, block):
    for idx in block:
        if idx < 0 or idx >= len(df_align):
            return False
        if str(df_align.at[idx, "坏点标记"]) == "坏点":
            return False
        if str(df_align.at[idx, "对齐状态"]) != "对齐成功":
            return False
    return True


def _triplet_structure_cost(times5, score_times5):
    costs = 0.0
    for a in range(4):
        real_gap = times5[a + 1] - times5[a]
        score_gap = score_times5[a + 1] - score_times5[a]
        if not (np.isfinite(real_gap) and np.isfinite(score_gap) and real_gap > 0 and score_gap > 0):
            return np.inf
        costs += abs(math.log(real_gap / score_gap))
    return costs


def _triplet_target_time(left_t, right_t, left_s, right_s, s):
    if not (np.isfinite(left_t) and np.isfinite(right_t) and np.isfinite(left_s) and np.isfinite(right_s) and right_t > left_t and right_s > left_s):
        return np.nan
    ratio = (s - left_s) / max(right_s - left_s, 1e-9)
    ratio = min(max(ratio, 0.0), 1.0)
    return left_t + (right_t - left_t) * ratio


def _build_triplet_candidates(df_align, df_on, block, left_anchor, right_anchor):
    left_on = safe_float(df_align.at[left_anchor, "DP2候选onset_idx"])
    right_on = safe_float(df_align.at[right_anchor, "DP2候选onset_idx"])
    if not np.isfinite(left_on) or not np.isfinite(right_on):
        return None
    left_on, right_on = int(left_on), int(right_on)
    if right_on - left_on <= 3:
        return None

    left_t = float(df_align.at[left_anchor, "红点时间(s)"])
    right_t = float(df_align.at[right_anchor, "红点时间(s)"])
    left_s = float(df_align.at[left_anchor, "谱面时间_原始(s)"])
    right_s = float(df_align.at[right_anchor, "谱面时间_原始(s)"])
    if not (np.isfinite(left_t) and np.isfinite(right_t) and np.isfinite(left_s) and np.isfinite(right_s)):
        return None

    all_on = list(range(left_on + 1, right_on))
    block_cands = []
    for idx in block:
        score_f = float(df_align.at[idx, "谱面频率(Hz)"])
        score_s = float(df_align.at[idx, "谱面时间_原始(s)"])
        target_t = _triplet_target_time(left_t, right_t, left_s, right_s, score_s)
        cur = []
        for on_idx in all_on:
            on_t = float(df_on.at[on_idx, ONSET_T_COL])
            on_f = float(df_on.at[on_idx, ONSET_F0_COL])
            abs_c = abs(cents_error(score_f, on_f))
            if abs_c > TRIPLET_PITCH_CENTS_MAX:
                continue
            cur.append({
                "onset_idx": int(on_idx),
                "time": on_t,
                "f_raw": on_f,
                "abs_c": abs_c,
                "target_err": abs(on_t - target_t) if np.isfinite(target_t) else 0.0,
            })
        if len(cur) == 0:
            return None
        cur.sort(key=lambda x: (x["target_err"], x["abs_c"], x["time"]))
        block_cands.append(cur[:TRIPLET_MAX_CANDS_PER_NOTE])
    return block_cands


def _triplet_total_cost(df_align, block, left_anchor, right_anchor, triplet_times, triplet_pitch_abs):
    left_t = float(df_align.at[left_anchor, "红点时间(s)"])
    right_t = float(df_align.at[right_anchor, "红点时间(s)"])
    left_s = float(df_align.at[left_anchor, "谱面时间_原始(s)"])
    right_s = float(df_align.at[right_anchor, "谱面时间_原始(s)"])
    score_times = [left_s] + [float(df_align.at[idx, "谱面时间_原始(s)"]) for idx in block] + [right_s]
    real_times = [left_t] + list(triplet_times) + [right_t]

    dur_cost = _triplet_structure_cost(real_times, score_times)
    if not np.isfinite(dur_cost):
        return np.inf

    pitch_cost = sum(float(c) for c in triplet_pitch_abs) / 100.0
    target_cost = 0.0
    for idx, t in zip(block, triplet_times):
        target_t = _triplet_target_time(left_t, right_t, left_s, right_s, float(df_align.at[idx, "谱面时间_原始(s)"]))
        if np.isfinite(target_t):
            target_cost += abs(t - target_t) / max(right_t - left_t, 1e-9)

    return TRIPLET_DUR_WEIGHT * dur_cost + TRIPLET_PITCH_WEIGHT * pitch_cost + TRIPLET_TARGET_WEIGHT * target_cost


def repair_triplet_from_duration_f(df_align, df_on):
    repaired = 0
    touched = set()

    for i in range(1, len(df_align) - 1):
        if i in touched or (i - 1) in touched or (i + 1) in touched:
            continue
        if str(df_align.at[i, "时值等级"]) != "F":
            continue

        block = [i - 1, i, i + 1]
        if not _triplet_block_valid(df_align, block):
            continue

        left_anchor = nearest_success_left(df_align, block[0])
        right_anchor = nearest_success_right(df_align, block[-1])
        if left_anchor is None or right_anchor is None:
            continue
        if left_anchor in block or right_anchor in block:
            continue

        cand_lists = _build_triplet_candidates(df_align, df_on, block, left_anchor, right_anchor)
        if cand_lists is None:
            continue

        old_times = [safe_float(df_align.at[idx, "红点时间(s)"]) for idx in block]
        old_abs = []
        ok_old = True
        for idx in block:
            score_f = float(df_align.at[idx, "谱面频率(Hz)"])
            used_f = safe_float(df_align.at[idx, "红点频率_used(Hz)"])
            if not (np.isfinite(used_f) and used_f > 0):
                ok_old = False
                break
            old_abs.append(abs(cents_error(score_f, used_f)))
        if not ok_old or not (old_times[0] < old_times[1] < old_times[2]):
            continue

        old_cost = _triplet_total_cost(df_align, block, left_anchor, right_anchor, old_times, old_abs)
        if not np.isfinite(old_cost):
            continue

        best = None
        for a, b, c in itertools.product(cand_lists[0], cand_lists[1], cand_lists[2]):
            if not (a["onset_idx"] < b["onset_idx"] < c["onset_idx"]):
                continue
            if not (a["time"] < b["time"] < c["time"]):
                continue
            new_times = [a["time"], b["time"], c["time"]]
            new_abs = [a["abs_c"], b["abs_c"], c["abs_c"]]
            total = _triplet_total_cost(df_align, block, left_anchor, right_anchor, new_times, new_abs)
            if not np.isfinite(total):
                continue
            cand = (total, a, b, c)
            if best is None or cand[0] < best[0]:
                best = cand

        if best is None:
            continue

        new_total = best[0]
        if old_cost - new_total >= TRIPLET_ACCEPT_IMPROVE:
            for idx, cand in zip(block, [best[1], best[2], best[3]]):
                overwrite_match_at(
                    df_align,
                    idx,
                    int(cand["onset_idx"]),
                    df_on,
                    note=f"节奏F三音重排: total {old_cost:.3f}->{new_total:.3f}",
                    source="节奏F三音重排",
                )
                touched.add(idx)
                repaired += 1

    return df_align, repaired


def build_metrics_table(df_align):
    total = len(df_align)
    status_ok = df_align["对齐状态"].astype(str).str.strip().eq("对齐成功")
    pitch_err_all = pd.to_numeric(
        df_align["音准误差(cents)"], errors="coerce"
    )

    unmatched_mask = ~status_ok
    wrong_pitch_mask = status_ok & pitch_err_all.abs().gt(100.0)
    # B = unmatched/conflict/deleted OR provisional match with |pitch error| > 100 cents.
    # This is reporting-only and must never modify df_align or its alignment statuses.
    b_error_mask = unmatched_mask | wrong_pitch_mask
    eval_valid_mask = ~b_error_mask
    df_valid = df_align[eval_valid_mask].copy()

    # Retain legacy status-only final-match fields for downstream compatibility.
    legacy_success = int(status_ok.sum())
    success = int(eval_valid_mask.sum())
    dropped = int(unmatched_mask.sum())
    wrong_pitch_count = int(wrong_pitch_mask.sum())
    b_error_count = int(b_error_mask.sum())
    suspicious_count = int((df_align["可疑点标记"].isin(["可疑点", "连带可疑点"])).sum()) if "可疑点标记" in df_align.columns else 0
    bad_count = int((df_align["坏点标记"] == "坏点").sum()) if "坏点标记" in df_align.columns else 0
    legacy_success_rate = legacy_success / total * 100.0 if total > 0 else np.nan
    b_success_rate = success / total * 100.0 if total > 0 else np.nan

    time_err = safe_numeric(df_valid["时间误差(ms)"]) if success > 0 else pd.Series(dtype=float)
    pitch_err = safe_numeric(df_valid["音准误差(cents)"]) if success > 0 else pd.Series(dtype=float)
    pitch_counts = df_valid["音准等级"].value_counts() if success > 0 else pd.Series(dtype=int)
    rhythm_counts = df_valid["节奏等级"].value_counts() if success > 0 else pd.Series(dtype=int)
    dur_counts = df_valid["时值等级"].value_counts() if success > 0 else pd.Series(dtype=int)
    repair_counts = df_valid["DP2修复来源"].fillna("").astype(str).value_counts() if success > 0 and "DP2修复来源" in df_valid.columns else pd.Series(dtype=int)
    score_shift_val = round(float(df_align["score_shift(s)"].iloc[0]) if total > 0 else 0.0, 6)

    return pd.DataFrame([{
        "总谱面音数": total,
        "可疑点总数(含连带)": suspicious_count,
        "坏点总数(删除范围内)": bad_count,
        "删除未匹配数": dropped,
        "错音高>100c数": wrong_pitch_count,
        "B类错误数(错音高或未匹配)": b_error_count,
        "B口径有效对应数": success,
        "B口径有效对应率(%)": round(b_success_rate, 2),
        "最终匹配成功数": legacy_success,
        "最终匹配率(%)": round(legacy_success_rate, 2),
        "score_shift(s)": score_shift_val,
        "平均时间误差(ms)": round(time_err.mean(), 2) if success > 0 else np.nan,
        "平均绝对时间误差(ms)": round(time_err.abs().mean(), 2) if success > 0 else np.nan,
        "平均音准误差(cents)": round(pitch_err.mean(), 2) if success > 0 else np.nan,
        "平均绝对音准误差(cents)": round(pitch_err.abs().mean(), 2) if success > 0 else np.nan,
        "音准S个数": int(pitch_counts.get("S", 0)),
        "音准A个数": int(pitch_counts.get("A", 0)),
        "音准B个数": int(pitch_counts.get("B", 0)),
        "音准C个数": int(pitch_counts.get("C", 0)),
        "音准D个数": int(pitch_counts.get("D", 0)),
        "音准E个数": int(pitch_counts.get("E", 0)),
        "音准F(成功点中)个数": int(pitch_counts.get("F", 0)),
        "节奏优秀个数": int(rhythm_counts.get("优秀", 0)),
        "节奏可接受个数": int(rhythm_counts.get("可接受", 0)),
        "节奏偏差较大个数": int(rhythm_counts.get("偏差较大", 0)),
        "时值A个数": int(dur_counts.get("A", 0)),
        "时值B个数": int(dur_counts.get("B", 0)),
        "时值C个数": int(dur_counts.get("C", 0)),
        "时值D个数": int(dur_counts.get("D", 0)),
        "时值F个数": int(dur_counts.get("F", 0)),
        "DP2坏区主体模块修复数": int(repair_counts.get("坏区主体模块修复", 0)),
        "DP2单点F修复数": int(repair_counts.get("单点F修复", 0)),
        "DP2节奏F三音重排音数": int(repair_counts.get("节奏F三音重排", 0)),
        "DP2坏区边界恢复数": int(repair_counts.get("坏区边界恢复", 0)),
    }])


def save_excel_with_highlight(df_align, excel_path):
    with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
        df_align.to_excel(writer, index=False, sheet_name="对齐结果")
        ws = writer.book["对齐结果"]
        fill_red = PatternFill(fill_type="solid", fgColor="FFC7CE")
        fill_yellow = PatternFill(fill_type="solid", fgColor="FFF2CC")
        fill_green = PatternFill(fill_type="solid", fgColor="E2F0D9")
        headers = [cell.value for cell in ws[1]]
        status_col = headers.index("对齐状态") + 1 if "对齐状态" in headers else None
        susp_col = headers.index("可疑点标记") + 1 if "可疑点标记" in headers else None
        source_col = headers.index("DP2修复来源") + 1 if "DP2修复来源" in headers else None
        for r in range(2, ws.max_row + 1):
            if status_col is not None:
                v = ws.cell(r, status_col).value
                if v != "对齐成功":
                    for c in range(1, ws.max_column + 1):
                        ws.cell(r, c).fill = fill_red
                    continue
            if source_col is not None:
                s = ws.cell(r, source_col).value
                if s not in [None, ""]:
                    for c in range(1, ws.max_column + 1):
                        ws.cell(r, c).fill = fill_green
                    continue
            if susp_col is not None:
                s = ws.cell(r, susp_col).value
                if s in ["可疑点", "连带可疑点"]:
                    for c in range(1, ws.max_column + 1):
                        if ws.cell(r, c).fill.fill_type is None:
                            ws.cell(r, c).fill = fill_yellow


def main():
    base_dir = os.path.abspath(
        os.environ.get("ERHU_WORK_DIR", os.path.dirname(os.path.abspath(__file__)))
    )
    df_sc, df_on, df_dp1, prefix, score_file, onset_file, dp1_file = load_data(base_dir)
    onset_times = _onset_time_array(df_on)
    df_align = ensure_columns(df_dp1.copy())
    df_align = build_onset_index_mapping(df_align, df_on, onset_times=onset_times)
    df_align = recalc_table_metrics(df_align)

    print(f"✅ score: {score_file}")
    print(f"✅ onset: {onset_file}")
    print(f"✅ DP1输入: {dp1_file}")
    print(f"🏷️ RUN_TAG = {RUN_TAG}")

    print("\n[阶段0] 坏区边界恢复...")
    df_align, n_restore = restore_bad_region_boundaries(df_align, df_on, onset_times=onset_times)
    df_align = build_onset_index_mapping(df_align, df_on, onset_times=onset_times)
    df_align = recalc_table_metrics(df_align)
    print(f"[阶段0] 恢复边界音数: {n_restore}")

    print("\n[阶段1] 坏区主体模块修复...")
    df_align, n_bad = repair_bad_region_bodies(df_align, df_sc, df_on)
    df_align = recalc_table_metrics(df_align)
    print(f"[阶段1] 修复主体音数: {n_bad}")

    print("\n[阶段2] 单点F与节奏F精修由 DP3 执行，DP2 不修改这些分支。")
    n_single = 0
    n_triplet = 0

    df_align = enforce_unique_onset_assignments(df_align)
    print("\n[阶段3] 统计并保存...")
    df_metrics = build_metrics_table(df_align)
    align_csv = os.path.join(base_dir, f"{prefix}_{RUN_TAG}_对齐结果_{METHOD_TAG_CN}.csv")
    metric_csv = os.path.join(base_dir, f"{prefix}_{RUN_TAG}_指标统计_{METHOD_TAG_CN}.csv")
    excel_xlsx = os.path.join(base_dir, f"{prefix}_{RUN_TAG}_对齐结果_{METHOD_TAG_CN}.xlsx")
    df_align.to_csv(align_csv, index=False, encoding="utf-8-sig")
    df_metrics.to_csv(metric_csv, index=False, encoding="utf-8-sig")
    save_excel_with_highlight(df_align, excel_xlsx)
    print(f"✅ 已保存: {align_csv}")
    print(f"✅ 已保存: {metric_csv}")
    print(f"✅ 已保存: {excel_xlsx}")


if __name__ == "__main__":
    main()
