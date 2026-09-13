import os
from itertools import combinations

import numpy as np
import pandas as pd
from openpyxl.styles import PatternFill

from common_utils import (
    cents_error,
    classify_duration_ratio_level,
    classify_pitch_level,
    classify_rhythm_level,
    find_first_with_suffix,
    freq_to_midi,
    midi_round,
    safe_float,
    safe_numeric,
)
from alignment_contract import (
    CURRENT_ONSET_INDEX_COL,
    ONSET_ID_COL,
    ORIGINAL_ONSET_INDEX_COL,
    choose_onset_file,
    ensure_onset_identity,
    finalize_onset_dataframe,
    onset_index_by_id,
    split_normal_and_weak_onsets,
)
from dp2_legacy import (
    ensure_columns as ensure_dp2_columns,
    enforce_unique_onset_assignments,
    recalc_table_metrics as recalc_dp2_metrics,
    repair_single_pitch_f,
    repair_triplet_from_duration_f,
)

DP3_IMPORTED_REFINEMENT_STAGES = ("单点F修复", "节奏F三音重排")

# =========================
# 全局配置
# =========================
RUN_TAG = "DP3"
METHOD_TAG_CN = "score主导"

# ===== 装饰音后卷补丁：只修正“前同音，后主音”的装饰音 =====
ENABLE_ORNAMENT_BACK_REFINE = True
ORNAMENT_PITCH_SOFT_FREE_CENTS = 25.0
ORNAMENT_PITCH_HARD_MAX_CENTS = 50.0
ORNAMENT_PITCH_SOFT_PENALTY_MAX = 0.20
ORNAMENT_PREV_DUR_WEIGHT = 1.00
ORNAMENT_GRACE_DUR_WEIGHT = 2.20
ORNAMENT_STAY_BIAS = 0.001
ORNAMENT_IMPROVE_MARGIN = 1e-9

# ===== 竞争区参数：只处理同音 =====
ENABLE_LOCAL_GROUP_REASSIGN = True
# A complete same-pitch run of at most eight score events is reassigned jointly.
# Natural runs longer than this keep the frozen four-note overlapping split so
# this extension does not alter the existing handling of very long tremolo-like
# passages.
MAX_GROUP_NOTES = 8
LEGACY_SPLIT_GROUP_NOTES = 4
LOCAL_CAND_RADIUS = 4
SAME_NOTE_CAND_CENTS = 75.0
LOCAL_KEEP_MARGIN = 1e-9

# 竞争区评分：只看节奏（n个音 -> n+1个区间）
LOCAL_INTERVAL_WEIGHT = 1.0

# ===== 同音竞争区防“重复红点当两个主音”硬约束 =====
# 只对连续“主音-主音”的同音转移生效；装饰音不受这个硬间隔约束。
# 典型用途：禁止 7.406s/7.414s、19.510s/19.518s 这类 8ms 假双音。
ENABLE_MAIN_SAME_NOTE_TINY_GAP_BLOCK = True
MIN_MAIN_SAME_NOTE_GAP_SEC = 0.050       # 主音-主音同音最小实际间隔：50ms
MIN_MAIN_SAME_NOTE_GAP_RATIO = 0.25      # 同时要求不小于理论推进时长的 25%

# ===== 前卷参数：先竞争，后前卷 =====
ENABLE_FORWARD_EARLY_REFINE = True
LOOKBACK_ONSET_NUM = 10
EARLY_MAX_CENTS = 80.0
EARLY_FREE_CENTS = 50.0
EARLY_SOFT_PENALTY_MAX = 0.20
EARLY_EARLY_BIAS = 0.01
EARLY_REWARD_WEIGHT = 0.60
EARLY_PENALTY_WEIGHT = 1.00

# A same-note local reassignment runs before the legacy forward pass.  Lock a
# reassigned entry only when it materially improves the pre-reassignment
# global timing residual without materially degrading pitch.
SAME_NOTE_FORWARD_LOCK_MIN_TIME_GAIN_SEC = 0.040
SAME_NOTE_FORWARD_LOCK_MAX_PITCH_DEGRADE_CENTS = 25.0

# ===== 后删坏点（只影响结果统计，不影响匹配过程）=====
POST_DROP_BAD_PITCH = True
POST_DROP_CENTS_TH = 100.0

RECOVERY_MAX_COST = 0.65
RECOVERY_MIN_ADVANTAGE = 0.15
RECOVERY_NORMAL_PITCH_MAX_CENTS = 70.0
RECOVERY_STRONG_PITCH_MAX_CENTS = 90.0
RECOVERY_ABSOLUTE_PITCH_MAX_CENTS = 100.0
RECOVERY_STRONG_SOURCES = {
    "energy+yellow",
    "energy+plateau",
    "energy+platform_transition",
}
RECOVERY_SOURCE_PENALTY = {
    "energy+yellow": 0.00,
    "energy+plateau": 0.00,
    "energy": 0.05,
    "plateau": 0.08,
    "yellow": 0.12,
    "platform_transition_weak": 0.15,
}

# A conflict inside a consecutive repeated-pitch run cannot be recovered one
# score event at a time: the already matched sibling may own the earlier red.
# Re-open the whole short run and solve an ordered one-to-one assignment.
ENABLE_REPEATED_PITCH_CONFLICT_REASSIGN = True
REPEATED_PITCH_CONFLICT_MAX_NOTES = 4
REPEATED_PITCH_CONFLICT_MAX_COST = 1.50

# With two reliable anchors, order already provides a finite search interval.
# Keep the interpolated time as a soft cost, but do not truncate that interval
# with the one-sided fixed-window ambiguity rule.
RECOVERY_BILATERAL_MIN_ADVANTAGE = 0.03

# Weak platform candidates are deliberately excluded from DP1/DP2.  DP3 can
# use one only to close a direct `recovered_unmatched -> unmatched` gap, after
# the normal conservative recovery has already run.
ENABLE_PLATFORM_WEAK_GAP_RESCUE = True
PLATFORM_WEAK_MIN_CONFIDENCE = 0.80
PLATFORM_WEAK_MAX_PITCH_CENTS = 70.0
PLATFORM_WEAK_MAX_TIME_WINDOW_MS = 250.0
PLATFORM_WEAK_AMBIGUITY_MARGIN = 0.05

# Additional weak sources remain outside DP1/DP2 and are admitted only through
# their own structural gates in DP3.
VOICED_REENTRY_WEAK_SOURCE = "voiced_reentry_weak"
GRADUAL_PLATFORM_WEAK_SOURCE = "gradual_platform_transition_weak"
EXISTING_STRONG_GRACE_SOURCE = "existing_strong_grace_candidate"
STRUCTURED_WEAK_MIN_CONFIDENCE = 0.80
STRUCTURED_WEAK_MAX_PITCH_CENTS = 70.0
STRUCTURED_WEAK_MAX_COST = 0.65
STRUCTURED_WEAK_AMBIGUITY_MARGIN = 0.05

# ===== 列名（score）=====
SCORE_TIME_COL = "累计时间(s)"
SCORE_F0_COL = "音高频率(Hz)"
SCORE_NAME_COL = "音名"
SCORE_DUR_COL = "持续时间(s)"
SCORE_BPM_COL = "BPM"
SCORE_MISC_COL = "连音线和其他信息"
SCORE_LEGATO_COL = "连奏信息"
SCORE_GRACE_COL = "装饰音判定"

# ===== 列名（onset）=====
ONSET_F0_COL = "note_f0(Hz)"
ONSET_T_COL = "note_time(s)"


def find_score_file(base_dir):
    return find_first_with_suffix(base_dir, "_score_final.csv", "*_score_final.csv")


def get_prefix_from_score(score_file):
    suffix = "_score_final.csv"
    if not score_file.endswith(suffix):
        raise RuntimeError(f"文件名不符合规则: {score_file}")
    return score_file[:-len(suffix)]


def find_onset_file(base_dir, prefix):
    return choose_onset_file(base_dir, prefix=prefix).name


def find_dp2_align_file(base_dir, prefix):
    target = f"{prefix}_DP2_对齐结果_{METHOD_TAG_CN}.csv"
    target_path = os.path.join(base_dir, target)
    if os.path.exists(target_path):
        return target
    cand = [f for f in os.listdir(base_dir) if f.endswith(f"_DP2_对齐结果_{METHOD_TAG_CN}.csv")]
    if len(cand) == 0:
        raise RuntimeError(f"未找到 {target}")
    if len(cand) > 1:
        print(f"⚠️ 未找到同名前缀的 {target}，检测到多个 DP2 结果，默认使用第一个：{cand[0]}")
    else:
        print(f"⚠️ 未找到同名前缀的 {target}，改用：{cand[0]}")
    return cand[0]


def observed_black_time_from_onsets(df_on_all):
    """Return the physical end marker before pitch-row filtering removes it."""
    if "pre_src" not in df_on_all.columns:
        return np.nan
    source = df_on_all["pre_src"].fillna("").astype(str).str.strip().str.lower()
    black_rows = df_on_all.loc[source.eq("black")]
    if black_rows.empty:
        return np.nan
    for column in ("pre_time(s)", ONSET_T_COL):
        if column not in black_rows.columns:
            continue
        values = safe_numeric(black_rows[column]).dropna()
        if len(values):
            value = float(values.iloc[-1])
            if np.isfinite(value):
                return value
    return np.nan


def load_data(score_csv, onset_csv):
    df_sc = pd.read_csv(score_csv)
    df_on_all = ensure_onset_identity(pd.read_csv(onset_csv))
    observed_black_time = observed_black_time_from_onsets(df_on_all)
    df_on, df_weak = split_normal_and_weak_onsets(df_on_all)

    for col in [SCORE_F0_COL]:
        if col not in df_sc.columns:
            raise RuntimeError(f"score 文件缺少列: {col}")
    for col in [ONSET_F0_COL, ONSET_T_COL]:
        if col not in df_on.columns:
            raise RuntimeError(f"onset 文件缺少列: {col}")

    df_sc[SCORE_F0_COL] = safe_numeric(df_sc[SCORE_F0_COL])
    df_on[ONSET_F0_COL] = safe_numeric(df_on[ONSET_F0_COL])
    df_on[ONSET_T_COL] = safe_numeric(df_on[ONSET_T_COL])
    if not df_weak.empty:
        df_weak[ONSET_F0_COL] = safe_numeric(df_weak[ONSET_F0_COL])
        df_weak[ONSET_T_COL] = safe_numeric(df_weak[ONSET_T_COL])

    if SCORE_TIME_COL in df_sc.columns:
        df_sc[SCORE_TIME_COL] = safe_numeric(df_sc[SCORE_TIME_COL])
    if SCORE_DUR_COL in df_sc.columns:
        df_sc[SCORE_DUR_COL] = safe_numeric(df_sc[SCORE_DUR_COL])
    if SCORE_BPM_COL in df_sc.columns:
        df_sc[SCORE_BPM_COL] = safe_numeric(df_sc[SCORE_BPM_COL])

    df_sc = df_sc[np.isfinite(df_sc[SCORE_F0_COL])].reset_index(drop=True)
    df_on = df_on[np.isfinite(df_on[ONSET_F0_COL]) & np.isfinite(df_on[ONSET_T_COL])].reset_index(drop=True)
    df_weak = df_weak[
        np.isfinite(df_weak[ONSET_F0_COL]) & np.isfinite(df_weak[ONSET_T_COL])
    ].reset_index(drop=True)
    df_on = finalize_onset_dataframe(df_on)
    df_weak = finalize_onset_dataframe(df_weak)
    return df_sc, df_on, df_weak, observed_black_time


def get_score_time_array(df_sc):
    if SCORE_TIME_COL not in df_sc.columns:
        return None
    arr = safe_numeric(df_sc[SCORE_TIME_COL]).to_numpy(dtype=float)
    if np.all(np.isfinite(arr)):
        return arr
    return None


def get_onset_time_array(df_on):
    return safe_numeric(df_on[ONSET_T_COL]).to_numpy(dtype=float)


def get_score_duration_array(df_sc):
    if SCORE_DUR_COL in df_sc.columns:
        dur = safe_numeric(df_sc[SCORE_DUR_COL]).to_numpy(dtype=float)
        if len(dur) > 0 and np.any(np.isfinite(dur)):
            return dur
    score_t = get_score_time_array(df_sc)
    if score_t is None:
        return None

    dur = np.full(len(score_t), np.nan, dtype=float)
    for i in range(len(score_t) - 1):
        d = score_t[i + 1] - score_t[i]
        dur[i] = d if d > 0 else np.nan

    if len(score_t) >= 2:
        dur[-1] = dur[-2]
    elif len(score_t) == 1:
        dur[-1] = np.nan
    return dur


def is_grace_note(sc_idx, df_sc):
    if SCORE_GRACE_COL not in df_sc.columns:
        return False
    val = df_sc.loc[sc_idx, SCORE_GRACE_COL]
    if pd.isna(val):
        return False
    return str(val).strip() == "装饰音"


def get_effective_prev_duration(prev_sc_idx, df_sc, score_dur):
    if score_dur is None:
        return np.nan
    if prev_sc_idx < 0 or prev_sc_idx >= len(score_dur):
        return np.nan

    if is_grace_note(prev_sc_idx, df_sc):
        if SCORE_BPM_COL in df_sc.columns and pd.notna(df_sc.loc[prev_sc_idx, SCORE_BPM_COL]):
            bpm = float(df_sc.loc[prev_sc_idx, SCORE_BPM_COL])
            if np.isfinite(bpm) and bpm > 0:
                return 6.0 / bpm
    return score_dur[prev_sc_idx]


def get_grace_target_duration(sc_idx, df_sc):
    if SCORE_BPM_COL not in df_sc.columns:
        return np.nan
    bpm = pd.to_numeric(pd.Series([df_sc.loc[sc_idx, SCORE_BPM_COL]]), errors="coerce").iloc[0]
    if np.isfinite(bpm) and bpm > 0:
        return 6.0 / bpm
    return np.nan


# ============================================================
# 从 DP2 恢复 pairs
# ============================================================
def load_pairs_from_dp2(dp2_csv, df_sc, df_on):
    df_dp2 = dp2_csv.copy() if isinstance(dp2_csv, pd.DataFrame) else pd.read_csv(dp2_csv)
    need_cols = ["红点时间(s)", "红点频率_raw(Hz)"]
    for c in need_cols:
        if c not in df_dp2.columns:
            raise RuntimeError(f"DP2 结果表缺少列: {c}")

    if len(df_dp2) != len(df_sc):
        raise RuntimeError(f"DP2 表行数 {len(df_dp2)} 与 score 行数 {len(df_sc)} 不一致")

    onset_t = pd.to_numeric(df_on[ONSET_T_COL], errors="coerce").to_numpy(dtype=float)
    onset_f = pd.to_numeric(df_on[ONSET_F0_COL], errors="coerce").to_numpy(dtype=float)

    used = set()
    pairs = []
    for sc_idx in range(len(df_dp2)):
        r = df_dp2.iloc[sc_idx]
        red_t = pd.to_numeric(pd.Series([r["红点时间(s)"]]), errors="coerce").iloc[0]
        red_f_raw = pd.to_numeric(pd.Series([r["红点频率_raw(Hz)"]]), errors="coerce").iloc[0]
        if not np.isfinite(red_t) or not np.isfinite(red_f_raw):
            continue
        stable_id = str(r.get(ONSET_ID_COL, "")).strip()
        if stable_id:
            mapped = onset_index_by_id(df_on, stable_id)
            if mapped is None:
                raise ValueError(f"DP2 引用了不存在的 onset_id: {stable_id}")
            on_idx = mapped
        else:
            # Compatibility path for historical DP2 files only.
            cand_idx = np.where(np.isclose(onset_t, red_t, atol=1e-9) & np.isclose(onset_f, red_f_raw, atol=1e-6))[0]
            if len(cand_idx) != 1:
                continue
            on_idx = int(cand_idx[0])
        if on_idx in used:
            continue
        used.add(on_idx)
        pairs.append((sc_idx, on_idx, "否", False))
    return pairs


# ============================================================
# 装饰音后卷补丁
# ============================================================
def log_ratio_error(actual_gap, target_gap):
    if not np.isfinite(actual_gap) or not np.isfinite(target_gap) or actual_gap <= 0 or target_gap <= 0:
        return np.inf
    return abs(np.log(actual_gap / target_gap))


def ornament_pitch_penalty(abs_c):
    if abs_c <= ORNAMENT_PITCH_SOFT_FREE_CENTS:
        return 0.0
    if abs_c <= ORNAMENT_PITCH_HARD_MAX_CENTS:
        frac = (abs_c - ORNAMENT_PITCH_SOFT_FREE_CENTS) / max(
            ORNAMENT_PITCH_HARD_MAX_CENTS - ORNAMENT_PITCH_SOFT_FREE_CENTS, 1e-9
        )
        return ORNAMENT_PITCH_SOFT_PENALTY_MAX * frac
    return np.inf


def ornament_back_refine(pairs, df_sc, df_on, sc_m, sc_f):
    if not ENABLE_ORNAMENT_BACK_REFINE or len(pairs) < 3:
        return pairs

    onset_t = get_onset_time_array(df_on)
    onset_f = df_on[ONSET_F0_COL].to_numpy(dtype=float)
    score_dur = get_score_duration_array(df_sc)

    new_pairs = pairs.copy()
    change_cnt = 0

    print("\n[阶段2] 装饰音后卷修正...")

    for k in range(1, len(new_pairs) - 1):
        prev_sc_idx, prev_on_idx, _, _ = new_pairs[k - 1]
        cur_sc_idx, cur_on_idx, _, _ = new_pairs[k]
        next_sc_idx, next_on_idx, _, _ = new_pairs[k + 1]

        # 条件1：当前音必须是装饰音；后一个必须是主音；前一个与当前必须同音
        if not is_grace_note(cur_sc_idx, df_sc):
            continue
        if is_grace_note(next_sc_idx, df_sc):
            continue
        if int(sc_m[prev_sc_idx]) != int(sc_m[cur_sc_idx]):
            continue

        # 只能在当前红点到后一个主音之前后卷
        if not (0 <= prev_on_idx < cur_on_idx < next_on_idx < len(df_on)):
            continue

        prev_target = get_effective_prev_duration(prev_sc_idx, df_sc, score_dur)
        cur_target = get_grace_target_duration(cur_sc_idx, df_sc)
        if not np.isfinite(prev_target) or not np.isfinite(cur_target):
            continue

        old_prev_gap = onset_t[cur_on_idx] - onset_t[prev_on_idx]
        old_cur_gap = onset_t[next_on_idx] - onset_t[cur_on_idx]
        old_prev_err = log_ratio_error(old_prev_gap, prev_target)
        old_cur_err = log_ratio_error(old_cur_gap, cur_target)
        if not np.isfinite(old_prev_err) or not np.isfinite(old_cur_err):
            continue

        old_score = (
            ORNAMENT_PREV_DUR_WEIGHT * old_prev_err +
            ORNAMENT_GRACE_DUR_WEIGHT * old_cur_err
        )

        best_on_idx = cur_on_idx
        best_score = old_score

        for cand in range(cur_on_idx, next_on_idx):
            if cand <= prev_on_idx or cand >= next_on_idx:
                continue

            abs_c = abs(cents_error(sc_f[cur_sc_idx], onset_f[cand]))
            p_pen = ornament_pitch_penalty(abs_c)
            if not np.isfinite(p_pen):
                continue

            new_prev_gap = onset_t[cand] - onset_t[prev_on_idx]
            new_cur_gap = onset_t[next_on_idx] - onset_t[cand]
            prev_err = log_ratio_error(new_prev_gap, prev_target)
            cur_err = log_ratio_error(new_cur_gap, cur_target)
            if not np.isfinite(prev_err) or not np.isfinite(cur_err):
                continue

            score = (
                ORNAMENT_PREV_DUR_WEIGHT * prev_err +
                ORNAMENT_GRACE_DUR_WEIGHT * cur_err +
                p_pen +
                ORNAMENT_STAY_BIAS * abs(cand - cur_on_idx)
            )

            # 条件2+3：只有整体更好才动
            if score + ORNAMENT_IMPROVE_MARGIN < best_score:
                best_score = score
                best_on_idx = cand

        if best_on_idx != cur_on_idx:
            new_pairs[k] = (cur_sc_idx, best_on_idx, "否", False)
            change_cnt += 1
            print(
                f"  ✅ 装饰音 sc={cur_sc_idx} 后卷: {cur_on_idx} -> {best_on_idx} | "
                f"prev_target={prev_target:.4f}s grace_target={cur_target:.4f}s"
            )

    if change_cnt == 0:
        print("  - 未发现满足条件且能带来改进的装饰音后卷")
    else:
        print(f"  ✅ 共修正 {change_cnt} 个装饰音位置")

    return new_pairs


# ============================================================
# 同音竞争区：只认 diff==0
# ============================================================
def has_same_note_competition(pair_a, pair_b, sc_m):
    sc_idx_a, _, _, _ = pair_a
    sc_idx_b, _, _, _ = pair_b
    return int(sc_m[sc_idx_a]) == int(sc_m[sc_idx_b])


def detect_same_note_groups(pairs, sc_m, max_group_notes=MAX_GROUP_NOTES):
    groups = []
    n = len(pairs)
    if n < 2:
        return groups

    i = 0
    while i < n - 1:
        if not has_same_note_competition(pairs[i], pairs[i + 1], sc_m):
            i += 1
            continue

        start = i
        j = i + 1
        while j < n - 1 and has_same_note_competition(pairs[j], pairs[j + 1], sc_m):
            j += 1

        run_len = j - start + 1
        if run_len <= max_group_notes:
            groups.append((start, j))
        else:
            cur = start
            while cur < j:
                end = min(cur + LEGACY_SPLIT_GROUP_NOTES - 1, j)
                if end - cur + 1 >= 2:
                    groups.append((cur, end))
                if end == j:
                    break
                cur = end - 1
        i = j + 1
    return groups


def build_local_candidate_lists(
    pairs,
    k_start,
    k_end,
    onset_len,
    right_on=None,
    use_full_anchor_interval=False,
):
    if k_start - 1 < 0:
        return None
    left_on = pairs[k_start - 1][1]
    if right_on is None:
        if k_end + 1 >= len(pairs):
            return None
        right_on = pairs[k_end + 1][1]
    if left_on >= right_on:
        return None

    cand_lists = []
    for k in range(k_start, k_end + 1):
        if use_full_anchor_interval:
            lo = left_on + 1
            hi = right_on - 1
        else:
            base_on = pairs[k][1]
            lo = max(left_on + 1, base_on - LOCAL_CAND_RADIUS)
            hi = min(right_on - 1, base_on + LOCAL_CAND_RADIUS)
        cand = list(range(lo, hi + 1))
        cand = sorted(set([x for x in cand if 0 <= x < onset_len]))
        if len(cand) == 0:
            return None
        cand_lists.append(cand)
    return cand_lists


def interval_ratio_error(actual_gap, score_gap):
    if not np.isfinite(actual_gap) or not np.isfinite(score_gap) or actual_gap <= 0 or score_gap <= 0:
        return np.inf
    return abs(np.log(actual_gap / score_gap))


def is_main_note(sc_idx, df_sc):
    """主音判定：非装饰音即视为主音。"""
    return not is_grace_note(sc_idx, df_sc)


def main_same_note_min_gap(prev_sc_idx, cur_sc_idx, df_sc, score_t, sc_m):
    """
    连续同音主音的最小允许间隔。
    若不是“主音-主音同音”，返回 0，不做限制。
    """
    if not ENABLE_MAIN_SAME_NOTE_TINY_GAP_BLOCK:
        return 0.0
    if prev_sc_idx is None or cur_sc_idx is None:
        return 0.0
    if prev_sc_idx < 0 or cur_sc_idx < 0:
        return 0.0
    if prev_sc_idx >= len(df_sc) or cur_sc_idx >= len(df_sc):
        return 0.0
    if not (is_main_note(prev_sc_idx, df_sc) and is_main_note(cur_sc_idx, df_sc)):
        return 0.0
    if int(sc_m[prev_sc_idx]) != int(sc_m[cur_sc_idx]):
        return 0.0

    score_gap = np.nan
    if score_t is not None and prev_sc_idx < len(score_t) and cur_sc_idx < len(score_t):
        score_gap = float(score_t[cur_sc_idx] - score_t[prev_sc_idx])

    if np.isfinite(score_gap) and score_gap > 0:
        return max(MIN_MAIN_SAME_NOTE_GAP_SEC, MIN_MAIN_SAME_NOTE_GAP_RATIO * score_gap)
    return MIN_MAIN_SAME_NOTE_GAP_SEC


def main_same_note_gap_ok(prev_sc_idx, cur_sc_idx, prev_on_idx, cur_on_idx, df_sc, score_t, onset_t, sc_m):
    """
    判断连续同音主音候选间隔是否合法。
    用于同音竞争区 DP 搜索与 cost 评估，防止把同一物理起音附近的重复红点分给两个主音。
    """
    min_gap = main_same_note_min_gap(prev_sc_idx, cur_sc_idx, df_sc, score_t, sc_m)
    if min_gap <= 0:
        return True
    if not (0 <= prev_on_idx < len(onset_t) and 0 <= cur_on_idx < len(onset_t)):
        return False
    actual_gap = float(onset_t[cur_on_idx] - onset_t[prev_on_idx])
    return np.isfinite(actual_gap) and actual_gap >= min_gap


def local_transition_score_gap(prev_sc_idx, cur_sc_idx, score_t, df_sc):
    """
    局部同音竞争评分使用的谱面间隔。
    对装饰音->主音这类累计时间相同的转移，不能直接用 score_t 差值 0，
    否则竞争区会全部变成 inf，导致无法修复。
    若 prev 是装饰音，则使用 6/BPM 作为装饰音目标时值。
    """
    if prev_sc_idx is None or cur_sc_idx is None:
        return np.nan
    if score_t is None:
        return np.nan
    if prev_sc_idx < 0 or cur_sc_idx < 0 or prev_sc_idx >= len(score_t) or cur_sc_idx >= len(score_t):
        return np.nan

    raw_gap = float(score_t[cur_sc_idx] - score_t[prev_sc_idx])
    if np.isfinite(raw_gap) and raw_gap > 1e-9:
        return raw_gap

    # 装饰音通常与目标主音共享累计时间；这里用 6/BPM 给它一个可比较的理论间隔。
    if is_grace_note(prev_sc_idx, df_sc):
        g = get_grace_target_duration(prev_sc_idx, df_sc)
        if np.isfinite(g) and g > 0:
            return float(g)

    # 兜底：如果持续时间列有正值，就用前一音持续时间。
    if SCORE_DUR_COL in df_sc.columns and pd.notna(df_sc.loc[prev_sc_idx, SCORE_DUR_COL]):
        d = float(df_sc.loc[prev_sc_idx, SCORE_DUR_COL])
        if np.isfinite(d) and d > 0:
            return d

    return np.nan


def candidate_pitch_ok(sc_idx, on_idx, sc_f, on_f, max_cents=SAME_NOTE_CAND_CENTS):
    abs_c = abs(cents_error(sc_f[sc_idx], on_f[on_idx]))
    return abs_c <= max_cents


def estimate_score_shift_from_pairs(pairs, score_t, onset_t):
    """
    用当前 pairs 粗估 score->onset 的全局平移。
    尾部虚拟 black 右锚点需要落在 onset 时间轴上，
    因此这里用 red_t - score_t 的中位数，抗尾部错配。
    """
    vals = []
    for sc_idx, on_idx, _, _ in pairs:
        if 0 <= sc_idx < len(score_t) and 0 <= on_idx < len(onset_t):
            v = onset_t[on_idx] - score_t[sc_idx]
            if np.isfinite(v):
                vals.append(float(v))
    if not vals:
        return 0.0
    return float(np.median(vals))


def eval_same_note_group_cost(group_pairs, left_pair, right_pair, score_t, onset_t,
                              right_score_time=None, right_on_time=None,
                              df_sc=None, sc_m=None):
    if len(group_pairs) == 0:
        return np.inf

    score_indices_local = [left_pair[0]] + [x[0] for x in group_pairs]
    onset_indices_local = [left_pair[1]] + [x[1] for x in group_pairs]
    score_times_local = [score_t[left_pair[0]]] + [score_t[x[0]] for x in group_pairs]
    onset_times_local = [onset_t[left_pair[1]]] + [onset_t[x[1]] for x in group_pairs]

    if right_pair is not None:
        score_indices_local.append(right_pair[0])
        onset_indices_local.append(right_pair[1])
        score_times_local.append(score_t[right_pair[0]])
        onset_times_local.append(onset_t[right_pair[1]])
    else:
        if right_score_time is None or right_on_time is None:
            return np.inf
        score_indices_local.append(None)
        onset_indices_local.append(None)
        score_times_local.append(float(right_score_time))
        onset_times_local.append(float(right_on_time))

    total = 0.0
    for i in range(len(score_times_local) - 1):
        # 同音主音硬间隔约束：禁止把 8ms/20ms 这类重复红点当成两个主音。
        if df_sc is not None and sc_m is not None and score_indices_local[i] is not None and score_indices_local[i + 1] is not None:
            if not main_same_note_gap_ok(
                score_indices_local[i],
                score_indices_local[i + 1],
                onset_indices_local[i],
                onset_indices_local[i + 1],
                df_sc, score_t, onset_t, sc_m
            ):
                return np.inf

        if df_sc is not None and score_indices_local[i] is not None and score_indices_local[i + 1] is not None:
            sg = local_transition_score_gap(score_indices_local[i], score_indices_local[i + 1], score_t, df_sc)
        else:
            sg = score_times_local[i + 1] - score_times_local[i]
        og = onset_times_local[i + 1] - onset_times_local[i]
        err = interval_ratio_error(og, sg)
        if not np.isfinite(err):
            return np.inf
        total += err
    return LOCAL_INTERVAL_WEIGHT * total / max(len(score_times_local) - 1, 1)


def resolve_one_same_note_group(
    pairs,
    k_start,
    k_end,
    df_sc,
    df_on,
    sc_m,
    sc_f,
    observed_black_time=np.nan,
):
    """
    同音竞争区重排。
    普通区间：用左右真实 score 音作为锚点。
    尾部区间：如果 k_end 已经是最后一个 score 音，则用“平移后的谱面终点 black”作为虚拟右锚点，
             这样最后一串同音不会因为没有 k_end+1 而被跳过。
    """
    if k_start - 1 < 0:
        return pairs

    score_t = get_score_time_array(df_sc)
    score_dur = get_score_duration_array(df_sc)
    onset_t = get_onset_time_array(df_on)
    on_f = df_on[ONSET_F0_COL].to_numpy(dtype=float)

    if score_t is None or score_dur is None or onset_t is None:
        return pairs

    left_pair = pairs[k_start - 1]
    group_pairs = pairs[k_start:k_end + 1]
    group_len = len(group_pairs)

    if group_len < 2:
        return pairs

    is_tail_group = (k_end + 1 >= len(pairs))

    if is_tail_group:
        last_sc_idx = group_pairs[-1][0]
        if last_sc_idx < 0 or last_sc_idx >= len(score_t):
            return pairs
        if not np.isfinite(score_t[last_sc_idx]) or not np.isfinite(score_dur[last_sc_idx]):
            return pairs

        right_score_time = float(score_t[last_sc_idx] + score_dur[last_sc_idx])
        score_shift = estimate_score_shift_from_pairs(pairs, score_t, onset_t)
        if np.isfinite(observed_black_time):
            right_on_time = float(observed_black_time)
            black_source = "observed"
        else:
            right_on_time = float(right_score_time + score_shift)
            black_source = "fallback_shifted_score"
        right_pair = None

        right_on_bound = int(np.searchsorted(onset_t, right_on_time, side="left"))
        right_on_bound = max(right_on_bound, left_pair[1] + 1)
        right_on_bound = min(right_on_bound, len(df_on))
    else:
        right_pair = pairs[k_end + 1]
        right_score_time = None
        right_on_time = None
        right_on_bound = right_pair[1]

    cand_lists = build_local_candidate_lists(
        pairs,
        k_start,
        k_end,
        len(df_on),
        right_on=right_on_bound,
        use_full_anchor_interval=(
            group_len > LEGACY_SPLIT_GROUP_NOTES
            or (is_tail_group and np.isfinite(observed_black_time))
        ),
    )
    if cand_lists is None:
        return pairs

    old_cost = eval_same_note_group_cost(
        group_pairs, left_pair, right_pair, score_t, onset_t,
        right_score_time=right_score_time,
        right_on_time=right_on_time,
        df_sc=df_sc,
        sc_m=sc_m,
    )

    dp = []
    prev_ptr = []

    first_sc_idx = group_pairs[0][0]
    dp0 = []
    ptr0 = []
    for on_idx in cand_lists[0]:
        if on_idx <= left_pair[1] or on_idx >= right_on_bound:
            dp0.append(np.inf)
            ptr0.append(None)
            continue
        if not candidate_pitch_ok(first_sc_idx, on_idx, sc_f, on_f):
            dp0.append(np.inf)
            ptr0.append(None)
            continue
        if not main_same_note_gap_ok(
            left_pair[0], first_sc_idx, left_pair[1], on_idx,
            df_sc, score_t, onset_t, sc_m
        ):
            dp0.append(np.inf)
            ptr0.append(None)
            continue
        err = interval_ratio_error(
            onset_t[on_idx] - onset_t[left_pair[1]],
            local_transition_score_gap(left_pair[0], first_sc_idx, score_t, df_sc)
        )
        dp0.append(err if np.isfinite(err) else np.inf)
        ptr0.append(None)

    dp.append(dp0)
    prev_ptr.append(ptr0)

    for layer in range(1, group_len):
        sc_idx = group_pairs[layer][0]
        prev_sc_idx = group_pairs[layer - 1][0]

        cur_dp = [np.inf] * len(cand_lists[layer])
        cur_ptr = [None] * len(cand_lists[layer])

        for j, on_idx in enumerate(cand_lists[layer]):
            if on_idx >= right_on_bound:
                continue
            if not candidate_pitch_ok(sc_idx, on_idx, sc_f, on_f):
                continue

            best = np.inf
            best_prev_j = None
            for pj, prev_on_idx in enumerate(cand_lists[layer - 1]):
                if dp[layer - 1][pj] == np.inf:
                    continue
                if prev_on_idx >= on_idx:
                    continue
                if not main_same_note_gap_ok(
                    prev_sc_idx, sc_idx, prev_on_idx, on_idx,
                    df_sc, score_t, onset_t, sc_m
                ):
                    continue
                err = interval_ratio_error(
                    onset_t[on_idx] - onset_t[prev_on_idx],
                    local_transition_score_gap(prev_sc_idx, sc_idx, score_t, df_sc)
                )
                if not np.isfinite(err):
                    continue
                cand = dp[layer - 1][pj] + err
                if cand < best:
                    best = cand
                    best_prev_j = pj

            cur_dp[j] = best
            cur_ptr[j] = best_prev_j

        dp.append(cur_dp)
        prev_ptr.append(cur_ptr)

    last_layer = group_len - 1
    last_sc_idx = group_pairs[-1][0]
    best_total = np.inf
    best_last_j = None

    for j, last_on_idx in enumerate(cand_lists[last_layer]):
        if dp[last_layer][j] == np.inf:
            continue
        if last_on_idx >= right_on_bound:
            continue

        if right_pair is not None:
            end_on_gap = onset_t[right_pair[1]] - onset_t[last_on_idx]
            end_score_gap = local_transition_score_gap(last_sc_idx, right_pair[0], score_t, df_sc)
        else:
            end_on_gap = right_on_time - onset_t[last_on_idx]
            end_score_gap = right_score_time - score_t[last_sc_idx]

        if right_pair is not None:
            if not main_same_note_gap_ok(
                last_sc_idx, right_pair[0], last_on_idx, right_pair[1],
                df_sc, score_t, onset_t, sc_m
            ):
                continue

        err = interval_ratio_error(end_on_gap, end_score_gap)
        if not np.isfinite(err):
            continue
        cand = dp[last_layer][j] + err
        if cand < best_total:
            best_total = cand
            best_last_j = j

    if best_last_j is None:
        return pairs

    chosen_on = [None] * group_len
    cur_j = best_last_j
    for layer in range(group_len - 1, -1, -1):
        chosen_on[layer] = cand_lists[layer][cur_j]
        cur_j = prev_ptr[layer][cur_j] if layer > 0 else None

    new_group = []
    for idx_local, on_idx in enumerate(chosen_on):
        sc_idx = group_pairs[idx_local][0]
        new_group.append((sc_idx, on_idx, "否", False))

    new_cost = eval_same_note_group_cost(
        new_group, left_pair, right_pair, score_t, onset_t,
        right_score_time=right_score_time,
        right_on_time=right_on_time,
        df_sc=df_sc,
        sc_m=sc_m,
    )

    if new_cost + LOCAL_KEEP_MARGIN < old_cost:
        out = pairs.copy()
        out[k_start:k_end + 1] = new_group
        old_on = [p[1] for p in group_pairs]
        new_on = [p[1] for p in new_group]
        old_t = [float(onset_t[x]) for x in old_on]
        new_t = [float(onset_t[x]) for x in new_on]
        if is_tail_group:
            print(
                f"  ✅ 尾部同音竞争区 [{k_start}, {k_end}] 使用 black 右锚点重排 | "
                f"score_black={right_score_time:.3f}s anchor={right_on_time:.3f}s "
                f"source={black_source}"
            )
        print(
            f"  ✅ 同音竞争区 [{k_start}, {k_end}] tiny-gap约束后重排 | "
            f"cost {old_cost:.4f}->{new_cost:.4f} | on {old_on}->{new_on} | "
            f"t {[round(x, 3) for x in old_t]}->{[round(x, 3) for x in new_t]}"
        )
        return out
    return pairs


def same_note_group_reassign(
    pairs,
    df_sc,
    df_on,
    sc_m,
    sc_f,
    *,
    return_locked=False,
    observed_black_time=np.nan,
):
    if not ENABLE_LOCAL_GROUP_REASSIGN:
        return (pairs, set()) if return_locked else pairs
    score_t = get_score_time_array(df_sc)
    if score_t is None:
        print("⚠️ score 缺少有效累计时间，跳过同音竞争区")
        return (pairs, set()) if return_locked else pairs

    groups = detect_same_note_groups(pairs, sc_m, max_group_notes=MAX_GROUP_NOTES)
    if not groups:
        print("同音竞争区：未检测到竞争区")
        return (pairs, set()) if return_locked else pairs

    print(f"同音竞争区：检测到 {len(groups)} 个竞争区 -> {groups}")
    new_pairs = pairs.copy()
    locked_score_indices = set()
    for k_start, k_end in groups:
        before = new_pairs[k_start:k_end + 1]
        new_pairs = resolve_one_same_note_group(
            new_pairs,
            k_start,
            k_end,
            df_sc,
            df_on,
            sc_m,
            sc_f,
            observed_black_time=observed_black_time,
        )
        after = new_pairs[k_start:k_end + 1]
        if before != after:
            # The final acceptance decision is made after a no-lock forward
            # pass in run_dp3.  Here we only record which entries the local
            # bidirectional search actually changed.
            locked_score_indices.update(
                int(new_pair[0])
                for old_pair, new_pair in zip(before, after)
                if int(old_pair[1]) != int(new_pair[1])
            )
            print(f"  ✅ 同音竞争区 [{k_start}, {k_end}] 已重排")
        else:
            print(f"  - 同音竞争区 [{k_start}, {k_end}] 保持原分配")
    if locked_score_indices:
        print(f"  🔒 同音重排锁定 score 音数: {len(locked_score_indices)}")
    return (new_pairs, locked_score_indices) if return_locked else new_pairs


# ============================================================
# 前卷：从前往后，评分而非死阈值
# ============================================================
def early_pitch_penalty(abs_c):
    if abs_c <= EARLY_FREE_CENTS:
        return 0.0
    if abs_c <= EARLY_MAX_CENTS:
        frac = (abs_c - EARLY_FREE_CENTS) / max(EARLY_MAX_CENTS - EARLY_FREE_CENTS, 1e-9)
        return EARLY_SOFT_PENALTY_MAX * frac
    return np.inf


def prev_gap_struct_error(prev_gap, prev_score_dur):
    if not np.isfinite(prev_gap) or not np.isfinite(prev_score_dur) or prev_gap <= 0 or prev_score_dur <= 0:
        return np.inf
    return abs(np.log(prev_gap / prev_score_dur))


def forward_early_refine(pairs, df_sc, df_on, sc_f, *, locked_score_indices=None):
    if not ENABLE_FORWARD_EARLY_REFINE or not pairs:
        return pairs

    onset_t = get_onset_time_array(df_on)
    on_f = df_on[ONSET_F0_COL].to_numpy(dtype=float)
    score_dur = get_score_duration_array(df_sc)

    new_pairs = pairs.copy()
    prev_final_on_idx = -1
    locked_score_indices = set(locked_score_indices or ())

    for k in range(len(new_pairs)):
        sc_idx, cur_on_idx, _, _ = new_pairs[k]
        if sc_idx in locked_score_indices:
            prev_final_on_idx = cur_on_idx
            continue
        next_on_idx = new_pairs[k + 1][1] if k + 1 < len(new_pairs) else len(df_on)
        left_bound = max(prev_final_on_idx + 1, cur_on_idx - LOOKBACK_ONSET_NUM)

        best_on_idx = cur_on_idx
        best_score = np.inf

        if k > 0:
            prev_sc_idx = new_pairs[k - 1][0]
            prev_on_idx = prev_final_on_idx
            prev_dur = get_effective_prev_duration(prev_sc_idx, df_sc, score_dur)
            old_prev_gap = onset_t[cur_on_idx] - onset_t[prev_on_idx]
            old_prev_err = prev_gap_struct_error(old_prev_gap, prev_dur)
        else:
            prev_sc_idx = None
            prev_on_idx = None
            prev_dur = np.nan
            old_prev_err = 0.0

        for cand in range(left_bound, cur_on_idx + 1):
            if cand <= prev_final_on_idx or cand >= next_on_idx:
                continue

            abs_c = abs(cents_error(sc_f[sc_idx], on_f[cand]))
            p_pen = early_pitch_penalty(abs_c)
            if not np.isfinite(p_pen):
                continue

            score = p_pen + EARLY_EARLY_BIAS * (cur_on_idx - cand)

            if k > 0:
                new_prev_gap = onset_t[cand] - onset_t[prev_on_idx]
                new_prev_err = prev_gap_struct_error(new_prev_gap, prev_dur)
                if not np.isfinite(new_prev_err):
                    continue

                delta = new_prev_err - old_prev_err
                if delta > 0:
                    score += EARLY_PENALTY_WEIGHT * delta
                else:
                    score -= EARLY_REWARD_WEIGHT * (-delta)

            if score < best_score:
                best_score = score
                best_on_idx = cand

        new_pairs[k] = (sc_idx, best_on_idx, "否", False)
        prev_final_on_idx = best_on_idx

    return new_pairs


def accept_same_note_forward_locks(
    baseline_pairs,
    protected_pairs,
    candidate_score_indices,
    df_sc,
    df_on,
    sc_f,
):
    """Keep only local locks that improve the completed legacy forward pass.

    The no-lock forward result is the regression baseline.  This deliberately
    evaluates the *actual* downstream outcome rather than a partial DP2 state,
    so a locally lower same-note cost cannot freeze a phrase-wide timing drift.
    """
    if not candidate_score_indices:
        return set()

    score_t = get_score_time_array(df_sc)
    if score_t is None:
        return set()
    onset_t = get_onset_time_array(df_on)
    on_f = df_on[ONSET_F0_COL].to_numpy(dtype=float)

    baseline_by_score = {int(sc_idx): int(on_idx) for sc_idx, on_idx, _, _ in baseline_pairs}
    protected_by_score = {int(sc_idx): int(on_idx) for sc_idx, on_idx, _, _ in protected_pairs}
    base_offsets = np.asarray(
        [onset_t[on_idx] - score_t[sc_idx] for sc_idx, on_idx in baseline_by_score.items()],
        dtype=float,
    )
    base_offsets = base_offsets[np.isfinite(base_offsets)]
    if not len(base_offsets):
        return set()
    reference_shift = float(np.mean(base_offsets))

    accepted = set()
    for sc_idx in candidate_score_indices:
        old_on_idx = baseline_by_score.get(int(sc_idx))
        new_on_idx = protected_by_score.get(int(sc_idx))
        if old_on_idx is None or new_on_idx is None or old_on_idx == new_on_idx:
            continue

        old_time_error = abs((onset_t[old_on_idx] - score_t[sc_idx]) - reference_shift)
        new_time_error = abs((onset_t[new_on_idx] - score_t[sc_idx]) - reference_shift)
        old_pitch_error = abs(cents_error(sc_f[sc_idx], on_f[old_on_idx]))
        new_pitch_error = abs(cents_error(sc_f[sc_idx], on_f[new_on_idx]))
        if (
            new_time_error + SAME_NOTE_FORWARD_LOCK_MIN_TIME_GAIN_SEC < old_time_error
            and new_pitch_error <= old_pitch_error + SAME_NOTE_FORWARD_LOCK_MAX_PITCH_DEGRADE_CENTS
        ):
            accepted.add(int(sc_idx))
            print(
                f"  same-note forward lock score[{sc_idx}]: "
                f"time {old_time_error * 1000:.1f}->{new_time_error * 1000:.1f} ms, "
                f"pitch {old_pitch_error:.1f}->{new_pitch_error:.1f} cents"
            )
    return accepted


def remove_duplicate_onset_pairs(pairs, onset_len):
    if not pairs:
        return pairs
    fixed = []
    used = set()
    prev_on = -1
    for sc_idx, on_idx, flag, use_fix in pairs:
        # Never manufacture prev_onset + 1. Conflicts remain unmatched.
        if not (0 <= int(on_idx) < onset_len):
            continue
        if int(on_idx) in used or int(on_idx) <= prev_on:
            continue
        fixed.append((sc_idx, int(on_idx), flag, use_fix))
        used.add(int(on_idx))
        prev_on = int(on_idx)
    return fixed


# ============================================================
# 输出
# ============================================================
def build_align_table(df_sc, df_on, pairs):
    temp_rows = []

    for sc_idx, on_idx, _, _ in pairs:
        sc = df_sc.loc[sc_idx]
        on = df_on.loc[on_idx]

        score_time_raw = float(sc[SCORE_TIME_COL]) if SCORE_TIME_COL in df_sc.columns and pd.notna(sc[SCORE_TIME_COL]) else np.nan
        score_f = float(sc[SCORE_F0_COL])

        red_t = float(on[ONSET_T_COL])
        red_f_raw = float(on[ONSET_F0_COL])
        red_f_used = red_f_raw

        raw_time_error_s = red_t - score_time_raw if np.isfinite(score_time_raw) else np.nan
        pitch_err_cents = cents_error(score_f, red_f_used)

        temp_rows.append({
            "sc_idx": sc_idx,
            "score_time_raw": score_time_raw,
            "score_f": score_f,
            "red_t": red_t,
            "red_f_raw": red_f_raw,
            "red_f_used": red_f_used,
            ONSET_ID_COL: on[ONSET_ID_COL],
            ORIGINAL_ONSET_INDEX_COL: on[ORIGINAL_ONSET_INDEX_COL],
            CURRENT_ONSET_INDEX_COL: on[CURRENT_ONSET_INDEX_COL],
            "low_oct_fix": "否",
            "octave_flag": "否",
            "raw_time_error_s": raw_time_error_s,
            "pitch_err_cents": pitch_err_cents,
            "音名": sc[SCORE_NAME_COL] if SCORE_NAME_COL in df_sc.columns else "",
            "持续时间(s)": float(sc[SCORE_DUR_COL]) if SCORE_DUR_COL in df_sc.columns and pd.notna(sc[SCORE_DUR_COL]) else np.nan,
            "累计时长(s)": score_time_raw,
            "连音线和其他信息": sc[SCORE_MISC_COL] if SCORE_MISC_COL in df_sc.columns else "",
            "连奏信息": sc[SCORE_LEGATO_COL] if SCORE_LEGATO_COL in df_sc.columns else "",
            "装饰音判定": sc[SCORE_GRACE_COL] if SCORE_GRACE_COL in df_sc.columns else "",
            "BPM": float(sc[SCORE_BPM_COL]) if SCORE_BPM_COL in df_sc.columns and pd.notna(sc[SCORE_BPM_COL]) else np.nan,
        })

    df_temp = pd.DataFrame(temp_rows)
    score_shift = float(df_temp["raw_time_error_s"].mean()) if len(df_temp) > 0 else 0.0

    rows = []
    for _, r in df_temp.iterrows():
        score_time_shifted = r["score_time_raw"] + score_shift if np.isfinite(r["score_time_raw"]) else np.nan
        time_err_ms = (r["red_t"] - score_time_shifted) * 1000.0 if np.isfinite(score_time_shifted) else np.nan
        pitch_err_cents = r["pitch_err_cents"]

        is_post_dropped = (
            POST_DROP_BAD_PITCH
            and np.isfinite(pitch_err_cents)
            and abs(float(pitch_err_cents)) > POST_DROP_CENTS_TH
        )

        align_status = "未匹配(后删)" if is_post_dropped else "对齐成功"
        mark_flag = "是" if is_post_dropped else ""
        pitch_level = "F" if is_post_dropped else classify_pitch_level(pitch_err_cents)
        rhythm_level = classify_rhythm_level(time_err_ms)

        rows.append({
            "_score_idx_internal": int(r["sc_idx"]),
            "谱面时间_原始(s)": r["score_time_raw"],
            "谱面时间_平移后(s)": score_time_shifted,
            "score_shift(s)": score_shift,
            "谱面频率(Hz)": r["score_f"],
            "红点时间(s)": r["red_t"],
            "红点频率_raw(Hz)": r["red_f_raw"],
            "红点频率_used(Hz)": r["red_f_used"],
            ONSET_ID_COL: r[ONSET_ID_COL],
            ORIGINAL_ONSET_INDEX_COL: r[ORIGINAL_ONSET_INDEX_COL],
            CURRENT_ONSET_INDEX_COL: r[CURRENT_ONSET_INDEX_COL],
            "低音八度修正": "否",
            "高八度等效匹配": "否",
            "时间误差(ms)": time_err_ms,
            "音准误差(cents)": pitch_err_cents,
            "窗口(s)": np.nan,
            "对齐状态": align_status,
            "后删标记": mark_flag,
            "音准等级": pitch_level,
            "节奏等级": rhythm_level,
            "音名": r["音名"],
            "持续时间(s)": r["持续时间(s)"],
            "累计时长(s)": r["累计时长(s)"],
            "连音线和其他信息": r["连音线和其他信息"],
            "连奏信息": r["连奏信息"],
            "装饰音判定": r["装饰音判定"],
            "BPM": r["BPM"],
        })

    return pd.DataFrame(rows)


def complete_alignment_rows(df_align, df_sc):
    """Preserve one output row per score note; missing pairs stay explicit."""
    by_score = {
        int(row["_score_idx_internal"]): row.to_dict()
        for _, row in df_align.iterrows()
        if np.isfinite(pd.to_numeric(pd.Series([row.get("_score_idx_internal")]), errors="coerce").iloc[0])
    }
    template_columns = list(df_align.columns)
    existing_shifts = pd.to_numeric(
        df_align.get("score_shift(s)", pd.Series(dtype=float)), errors="coerce"
    ).dropna()
    reference_shift = float(existing_shifts.median()) if len(existing_shifts) else 0.0
    rows = []
    for sc_idx in range(len(df_sc)):
        if sc_idx in by_score:
            rows.append(by_score[sc_idx])
            continue
        score = df_sc.loc[sc_idx]
        score_time_raw = safe_float(score.get(SCORE_TIME_COL, np.nan))
        score_time_shifted = (
            score_time_raw + reference_shift
            if np.isfinite(score_time_raw)
            else np.nan
        )
        row = {column: np.nan for column in template_columns}
        row.update({
            "_score_idx_internal": sc_idx,
            "谱面时间_原始(s)": score_time_raw,
            "谱面时间_平移后(s)": score_time_shifted,
            "谱面频率(Hz)": score.get(SCORE_F0_COL, np.nan),
            "score_shift(s)": reference_shift,
            "红点时间(s)": np.nan,
            "红点频率_raw(Hz)": np.nan,
            "红点频率_used(Hz)": np.nan,
            "时间误差(ms)": np.nan,
            "音准误差(cents)": np.nan,
            "窗口(s)": np.nan,
            "对齐状态": "conflict",
            "后删标记": "",
            "音准等级": "",
            "节奏等级": "",
            "音名": score.get(SCORE_NAME_COL, ""),
            "持续时间(s)": score.get(SCORE_DUR_COL, np.nan),
            "累计时长(s)": score.get(SCORE_TIME_COL, np.nan),
            "连音线和其他信息": score.get(SCORE_MISC_COL, ""),
            "连奏信息": score.get(SCORE_LEGATO_COL, ""),
            "装饰音判定": score.get(SCORE_GRACE_COL, ""),
            "BPM": score.get(SCORE_BPM_COL, np.nan),
            ONSET_ID_COL: "",
        })
        rows.append(row)
    out = pd.DataFrame(rows)
    return out.drop(columns=["_score_idx_internal"], errors="ignore")


def _recovery_status_priority(value):
    text = str(value).strip().lower()
    if "conflict" in text:
        return 0
    if "后删" in text or "deleted" in text:
        return 1
    if "未匹配" in text or "unmatched" in text:
        return 2
    return None


def _valid_matched_row(row):
    return str(row.get("对齐状态", "")).strip() in {"对齐成功", "recovered_unmatched"}


def _is_conservative_recovery_row(row):
    """Whether a row was recovered by the main conservative-recovery pass.

    Current DP3 normalizes successful main-pass recoveries to ``对齐成功`` so
    they count as ordinary correspondences. The weak bridge still needs to
    recognize that provenance without treating ordinary DP matches as bridge
    anchors.
    """
    status = str(row.get("对齐状态", "")).strip()
    if status == "recovered_unmatched":
        return True
    stage_tokens = {
        token.strip()
        for token in str(row.get("recovery_stage", "")).split("|")
        if token.strip()
    }
    return (
        "dp3_conservative_recovery" in stage_tokens
        and str(row.get("recovery_result", "")).strip() == "recovered"
        and _recovery_status_priority(row.get("recovery_original_status", ""))
        is not None
    )


def _nearest_recovery_anchors(df_align, index):
    left = right = None
    for candidate in range(index - 1, -1, -1):
        if _valid_matched_row(df_align.loc[candidate]) and np.isfinite(safe_float(df_align.at[candidate, "红点时间(s)"])):
            left = candidate
            break
    for candidate in range(index + 1, len(df_align)):
        if _valid_matched_row(df_align.loc[candidate]) and np.isfinite(safe_float(df_align.at[candidate, "红点时间(s)"])):
            right = candidate
            break
    return left, right


def _predict_recovery_time(df_align, index, left, right):
    score_time = safe_float(df_align.at[index, "谱面时间_原始(s)"])
    if not np.isfinite(score_time):
        return np.nan, 150.0, "fallback_150ms"
    local_gaps = []
    if left is not None:
        left_score = safe_float(df_align.at[left, "谱面时间_原始(s)"])
        left_time = safe_float(df_align.at[left, "红点时间(s)"])
        if np.isfinite(left_score) and score_time > left_score:
            local_gaps.append((score_time - left_score) * 1000.0)
    else:
        left_score = left_time = np.nan
    if right is not None:
        right_score = safe_float(df_align.at[right, "谱面时间_原始(s)"])
        right_time = safe_float(df_align.at[right, "红点时间(s)"])
        if np.isfinite(right_score) and right_score > score_time:
            local_gaps.append((right_score - score_time) * 1000.0)
    else:
        right_score = right_time = np.nan

    if left is not None and right is not None and right_score > left_score and right_time > left_time:
        ratio = (score_time - left_score) / (right_score - left_score)
        predicted = left_time + ratio * (right_time - left_time)
        source = "bilateral_linear_interpolation"
    elif left is not None and np.isfinite(left_score) and np.isfinite(left_time):
        predicted = left_time + (score_time - left_score)
        source = "left_anchor_projection"
    elif right is not None and np.isfinite(right_score) and np.isfinite(right_time):
        predicted = right_time - (right_score - score_time)
        source = "right_anchor_projection"
    else:
        return np.nan, 150.0, "fallback_150ms"

    if local_gaps:
        local_score_ioi_ms = min(local_gaps)
        window = min(250.0, max(80.0, 0.4 * local_score_ioi_ms))
    else:
        window = 150.0
        source += "|fallback_150ms"
    return float(predicted), float(window), source


def reassign_repeated_pitch_conflicts(df_align, df_sc, df_on):
    """Jointly reopen short repeated-pitch runs containing an unresolved row.

    The candidate set includes the red events currently owned by the run, but
    excludes every red owned by score events outside the run.  Candidate time
    and score order are both strict, so the result stays one-to-one.
    """
    if not ENABLE_REPEATED_PITCH_CONFLICT_REASSIGN or len(df_align) < 2:
        return df_align

    out = df_align.copy()
    recovery_defaults = {
        "alignment_status": out.get("对齐状态", pd.Series("", index=out.index)),
        "recovery_stage": "",
        "recovery_original_status": "",
        "recovery_onset_id": "",
        "recovery_cost": np.nan,
        "recovery_time_error_ms": np.nan,
        "recovery_pitch_error_cents": np.nan,
        "recovery_predicted_time": np.nan,
        "recovery_search_window_ms": np.nan,
        "recovery_candidate_count": 0,
        "recovery_prediction_source": "",
        "recovery_result": "",
    }
    for column, default in recovery_defaults.items():
        if column not in out.columns:
            out[column] = default
    score_t = get_score_time_array(df_sc)
    onset_t = get_onset_time_array(df_on)
    if score_t is None or onset_t is None:
        return out

    sc_f = safe_numeric(df_sc[SCORE_F0_COL]).to_numpy(dtype=float)
    sc_m = midi_round(sc_f)
    on_f = safe_numeric(df_on[ONSET_F0_COL]).to_numpy(dtype=float)
    onset_row_by_id = {
        str(row.get(ONSET_ID_COL, "")).strip(): (int(index), row)
        for index, row in df_on.iterrows()
        if str(row.get(ONSET_ID_COL, "")).strip()
    }

    repeated_runs = []
    start = 0
    while start < len(out):
        end = start
        while (
            end + 1 < len(out)
            and np.isfinite(sc_m[start])
            and np.isfinite(sc_m[end + 1])
            and sc_m[end + 1] == sc_m[start]
        ):
            end += 1
        if 2 <= end - start + 1 <= REPEATED_PITCH_CONFLICT_MAX_NOTES:
            statuses = [out.at[index, "对齐状态"] for index in range(start, end + 1)]
            if (
                any(_recovery_status_priority(status) is not None for status in statuses)
                and any(_valid_matched_row(out.loc[index]) for index in range(start, end + 1))
            ):
                repeated_runs.append((start, end))
        start = end + 1

    for start, end in repeated_runs:
        left, _ = _nearest_recovery_anchors(out, start)
        _, right = _nearest_recovery_anchors(out, end)
        if left is None or right is None:
            continue
        left_id = str(out.at[left, ONSET_ID_COL]).strip()
        right_id = str(out.at[right, ONSET_ID_COL]).strip()
        if left_id not in onset_row_by_id or right_id not in onset_row_by_id:
            continue
        left_on_idx = onset_row_by_id[left_id][0]
        right_on_idx = onset_row_by_id[right_id][0]
        if right_on_idx <= left_on_idx:
            continue

        run_indices = set(range(start, end + 1))
        blocked_ids = {
            str(row.get(ONSET_ID_COL, "")).strip()
            for index, row in out.iterrows()
            if index not in run_indices
            and _valid_matched_row(row)
            and str(row.get(ONSET_ID_COL, "")).strip()
        }
        target_midi = sc_m[start]
        physical_by_frame = {}
        for on_idx, onset in df_on.iterrows():
            onset_id = str(onset.get(ONSET_ID_COL, "")).strip()
            note_frame = safe_float(onset.get("note_frame"))
            onset_time = safe_float(onset.get(ONSET_T_COL))
            onset_freq = safe_float(onset.get(ONSET_F0_COL))
            source = str(onset.get("pre_src", "")).strip().lower()
            if (
                not onset_id
                or onset_id in blocked_ids
                or not np.isfinite(note_frame)
                or note_frame < 0
                or not (left_on_idx < int(on_idx) < right_on_idx)
                or not (np.isfinite(onset_time) and np.isfinite(onset_freq) and onset_freq > 0)
                or "black" in {part.strip() for part in source.split("+")}
                or midi_round(np.asarray([onset_freq], dtype=float))[0] != target_midi
                or not candidate_pitch_ok(start, int(on_idx), sc_f, on_f)
            ):
                continue
            frame_key = int(note_frame)
            item = (int(on_idx), onset_id, onset_time, onset_freq, source, onset)
            existing = physical_by_frame.get(frame_key)
            if existing is None:
                physical_by_frame[frame_key] = item
            else:
                old_penalty = RECOVERY_SOURCE_PENALTY.get(existing[4], 0.12)
                new_penalty = RECOVERY_SOURCE_PENALTY.get(source, 0.12)
                if (new_penalty, onset_id) < (old_penalty, existing[1]):
                    physical_by_frame[frame_key] = item

        candidates = sorted(physical_by_frame.values(), key=lambda item: (item[0], item[1]))
        group_len = end - start + 1
        if len(candidates) < group_len:
            continue

        left_pair = (left, left_on_idx, "否", False)
        right_pair = (right, right_on_idx, "否", False)
        scored = []
        for chosen in combinations(candidates, group_len):
            chosen_on = [item[0] for item in chosen]
            group_pairs = [
                (score_index, on_idx, "否", False)
                for score_index, on_idx in zip(range(start, end + 1), chosen_on)
            ]
            interval_cost = eval_same_note_group_cost(
                group_pairs,
                left_pair,
                right_pair,
                score_t,
                onset_t,
                df_sc=df_sc,
                sc_m=sc_m,
            )
            if not np.isfinite(interval_cost):
                continue
            pitch_cost = float(np.mean([
                abs(cents_error(sc_f[score_index], item[3])) / SAME_NOTE_CAND_CENTS
                for score_index, item in zip(range(start, end + 1), chosen)
            ]))
            source_cost = float(np.mean([
                RECOVERY_SOURCE_PENALTY.get(item[4], 0.12) for item in chosen
            ]))
            total_cost = interval_cost + 0.15 * pitch_cost + 0.05 * source_cost
            if total_cost <= REPEATED_PITCH_CONFLICT_MAX_COST:
                scored.append((total_cost, chosen))

        if not scored:
            continue
        scored.sort(key=lambda item: (item[0], tuple(x[2] for x in item[1])))
        best_cost, best = scored[0]
        anchor_span_ms = (
            safe_float(out.at[right, "红点时间(s)"])
            - safe_float(out.at[left, "红点时间(s)"])
        ) * 1000.0
        old_ids = [str(out.at[index, ONSET_ID_COL]).strip() for index in range(start, end + 1)]
        new_ids = [item[1] for item in best]

        for score_index, item in zip(range(start, end + 1), best):
            _on_idx, onset_id, onset_time, onset_freq, _source, onset_row = item
            original_status = str(out.at[score_index, "对齐状态"])
            pitch_error = cents_error(sc_f[score_index], onset_freq)
            shifted_score_time = safe_float(out.at[score_index, "谱面时间_平移后(s)"])
            official_time_error_ms = (
                (onset_time - shifted_score_time) * 1000.0
                if np.isfinite(shifted_score_time)
                else np.nan
            )
            out.at[score_index, "对齐状态"] = "对齐成功"
            out.at[score_index, "alignment_status"] = "对齐成功"
            out.at[score_index, ONSET_ID_COL] = onset_id
            out.at[score_index, ORIGINAL_ONSET_INDEX_COL] = onset_row.get(ORIGINAL_ONSET_INDEX_COL, np.nan)
            out.at[score_index, CURRENT_ONSET_INDEX_COL] = onset_row.get(CURRENT_ONSET_INDEX_COL, np.nan)
            out.at[score_index, "红点时间(s)"] = onset_time
            out.at[score_index, "红点频率_raw(Hz)"] = onset_freq
            out.at[score_index, "红点频率_used(Hz)"] = onset_freq
            out.at[score_index, "时间误差(ms)"] = official_time_error_ms
            out.at[score_index, "音准误差(cents)"] = pitch_error
            out.at[score_index, "音准等级"] = classify_pitch_level(pitch_error)
            out.at[score_index, "节奏等级"] = classify_rhythm_level(official_time_error_ms)
            out.at[score_index, "后删标记"] = ""
            out.at[score_index, "recovery_stage"] = "dp3_repeated_pitch_joint_reassignment"
            out.at[score_index, "recovery_original_status"] = original_status
            out.at[score_index, "recovery_onset_id"] = onset_id
            out.at[score_index, "recovery_cost"] = best_cost
            out.at[score_index, "recovery_time_error_ms"] = np.nan
            out.at[score_index, "recovery_pitch_error_cents"] = pitch_error
            out.at[score_index, "recovery_predicted_time"] = np.nan
            out.at[score_index, "recovery_search_window_ms"] = anchor_span_ms
            out.at[score_index, "recovery_candidate_count"] = len(candidates)
            out.at[score_index, "recovery_prediction_source"] = "ordered_joint_interval_optimization"
            out.at[score_index, "recovery_result"] = "joint_reassigned"

        print(
            f"[DP3 repeated-pitch recovery] score[{start}:{end}] "
            f"{old_ids} -> {new_ids}, cost={best_cost:.4f}"
        )

    return out


def recover_unmatched_from_free_onsets(df_align, df_sc, df_on):
    """One conservative pass over real, free onset_id candidates."""
    out = df_align.copy()
    recovery_defaults = {
        "alignment_status": out.get("对齐状态", pd.Series("", index=out.index)),
        "recovery_stage": "",
        "recovery_original_status": "",
        "recovery_onset_id": "",
        "recovery_cost": np.nan,
        "recovery_time_error_ms": np.nan,
        "recovery_pitch_error_cents": np.nan,
        "recovery_predicted_time": np.nan,
        "recovery_search_window_ms": np.nan,
        "recovery_candidate_count": 0,
        "recovery_prediction_source": "",
        "recovery_result": "",
    }
    for column, default in recovery_defaults.items():
        if column not in out.columns:
            out[column] = default

    used_ids = {
        str(row.get(ONSET_ID_COL, "")).strip()
        for _, row in out.iterrows()
        if _valid_matched_row(row) and str(row.get(ONSET_ID_COL, "")).strip()
    }
    physical_candidates = {}
    for _, onset in df_on.iterrows():
        onset_id = str(onset.get(ONSET_ID_COL, "")).strip()
        note_frame = safe_float(onset.get("note_frame"))
        onset_time = safe_float(onset.get(ONSET_T_COL))
        onset_freq = safe_float(onset.get(ONSET_F0_COL))
        source = str(onset.get("pre_src", "")).strip().lower()
        if not onset_id or onset_id in used_ids or not np.isfinite(note_frame) or note_frame < 0:
            continue
        if "black" in {part.strip() for part in source.split("+")}:
            continue
        if not (np.isfinite(onset_time) and np.isfinite(onset_freq) and onset_freq > 0):
            continue
        frame_key = int(note_frame)
        item = (onset_id, onset_time, onset_freq, source, onset)
        existing = physical_candidates.get(frame_key)
        if existing is None:
            physical_candidates[frame_key] = item
        else:
            old_penalty = RECOVERY_SOURCE_PENALTY.get(existing[3], 0.12)
            new_penalty = RECOVERY_SOURCE_PENALTY.get(source, 0.12)
            if (new_penalty, onset_id) < (old_penalty, existing[0]):
                physical_candidates[frame_key] = item
    # Multiple detector rows on the exact same red frame represent one
    # physical onset event. Keep one stable onset_id; never merge by fuzzy time.
    candidates = list(physical_candidates.values())

    targets = []
    for index in range(len(out)):
        priority = _recovery_status_priority(out.at[index, "对齐状态"])
        if priority is None:
            continue
        score_row = df_sc.loc[index]
        align_flag = str(score_row.get("参与起音对齐", "是")).strip().lower()
        score_freq = safe_float(score_row.get(SCORE_F0_COL))
        if align_flag in {"否", "no", "false", "0"} or not np.isfinite(score_freq) or score_freq <= 0:
            continue
        targets.append((priority, index))

    # Conflict/deleted/unmatched rows do not own their former candidate.
    # Release identity before recovery so output onset_id remains one-to-one.
    for _priority, index in targets:
        out.at[index, ONSET_ID_COL] = ""
        out.at[index, ORIGINAL_ONSET_INDEX_COL] = np.nan
        out.at[index, CURRENT_ONSET_INDEX_COL] = np.nan

    for _priority, index in sorted(targets):
        original_status = str(out.at[index, "对齐状态"])
        left, right = _nearest_recovery_anchors(out, index)
        if left is None and right is None:
            out.at[index, "recovery_result"] = "recovery_no_anchor"
            continue
        predicted, window_ms, prediction_source = _predict_recovery_time(out, index, left, right)
        if not np.isfinite(predicted):
            out.at[index, "recovery_result"] = "recovery_no_prediction"
            continue
        score_freq = safe_float(out.at[index, "谱面频率(Hz)"])
        bilateral_anchor_search = (
            left is not None
            and right is not None
            and safe_float(out.at[right, "红点时间(s)"])
            > safe_float(out.at[left, "红点时间(s)"])
        )
        anchor_window_ms = float(window_ms)
        anchor_prediction_source = prediction_source
        if bilateral_anchor_search:
            anchor_window_ms = max(
                anchor_window_ms,
                (
                    safe_float(out.at[right, "红点时间(s)"])
                    - safe_float(out.at[left, "红点时间(s)"])
                ) * 1000.0,
            )
            anchor_prediction_source += "|anchor_bounded_search"

        # Preserve the original short-window decision whenever it already has
        # a plausible physical onset. The newer anchor-bounded search remains
        # available as a fallback only when that local search is empty. This
        # keeps a nearby unique event from becoming ambiguous merely because a
        # distant, anchor-legal alternative was also admitted.
        local_qualified = []
        anchor_qualified = []
        for onset_id, onset_time, onset_freq, source, onset_row in candidates:
            if onset_id in used_ids:
                continue
            if left is not None and not (onset_time > safe_float(out.at[left, "红点时间(s)"])):
                continue
            if right is not None and not (onset_time < safe_float(out.at[right, "红点时间(s)"])):
                continue
            time_error_ms = (onset_time - predicted) * 1000.0
            pitch_error = cents_error(score_freq, onset_freq)
            abs_pitch = abs(pitch_error)
            pitch_limit = RECOVERY_STRONG_PITCH_MAX_CENTS if source in RECOVERY_STRONG_SOURCES else RECOVERY_NORMAL_PITCH_MAX_CENTS
            if abs_pitch > pitch_limit or abs_pitch > RECOVERY_ABSOLUTE_PITCH_MAX_CENTS:
                continue
            pitch_cost = abs_pitch / 100.0
            source_penalty = RECOVERY_SOURCE_PENALTY.get(source, 0.12)
            if abs(time_error_ms) <= window_ms:
                local_cost = (
                    0.55 * abs(time_error_ms) / window_ms
                    + 0.35 * pitch_cost
                    + 0.10 * source_penalty
                )
                if local_cost <= RECOVERY_MAX_COST:
                    local_qualified.append((local_cost, onset_id, onset_time, onset_freq, source, pitch_error, time_error_ms, onset_row))
            if bilateral_anchor_search:
                anchor_cost = (
                    0.55 * abs(time_error_ms) / anchor_window_ms
                    + 0.35 * pitch_cost
                    + 0.10 * source_penalty
                )
                if anchor_cost <= RECOVERY_MAX_COST:
                    anchor_qualified.append((anchor_cost, onset_id, onset_time, onset_freq, source, pitch_error, time_error_ms, onset_row))

        if local_qualified:
            qualified = local_qualified
            effective_window_ms = float(window_ms)
            prediction_source += "|local_window_primary"
            min_advantage = RECOVERY_MIN_ADVANTAGE
        elif bilateral_anchor_search:
            qualified = anchor_qualified
            effective_window_ms = anchor_window_ms
            prediction_source = anchor_prediction_source
            min_advantage = RECOVERY_BILATERAL_MIN_ADVANTAGE
        else:
            qualified = []
            effective_window_ms = float(window_ms)
            min_advantage = RECOVERY_MIN_ADVANTAGE

        qualified.sort(key=lambda item: (item[0], item[2], item[1]))
        out.at[index, "recovery_stage"] = "dp3_conservative_recovery"
        out.at[index, "recovery_original_status"] = original_status
        out.at[index, "recovery_predicted_time"] = predicted
        out.at[index, "recovery_search_window_ms"] = effective_window_ms
        out.at[index, "recovery_candidate_count"] = len(qualified)
        out.at[index, "recovery_prediction_source"] = prediction_source
        if not qualified:
            out.at[index, "recovery_result"] = "recovery_no_candidate"
            continue
        if len(qualified) > 1 and qualified[1][0] - qualified[0][0] < min_advantage:
            out.at[index, "recovery_result"] = "recovery_ambiguous"
            continue

        cost, onset_id, onset_time, onset_freq, _source, pitch_error, recovery_time_error_ms, onset_row = qualified[0]
        shifted_score_time = safe_float(out.at[index, "谱面时间_平移后(s)"])
        official_time_error_ms = (
            (onset_time - shifted_score_time) * 1000.0
            if np.isfinite(shifted_score_time)
            else np.nan
        )
        out.at[index, "对齐状态"] = "对齐成功"
        out.at[index, "alignment_status"] = "对齐成功"
        out.at[index, ONSET_ID_COL] = onset_id
        out.at[index, ORIGINAL_ONSET_INDEX_COL] = onset_row.get(ORIGINAL_ONSET_INDEX_COL, np.nan)
        out.at[index, CURRENT_ONSET_INDEX_COL] = onset_row.get(CURRENT_ONSET_INDEX_COL, np.nan)
        out.at[index, "红点时间(s)"] = onset_time
        out.at[index, "红点频率_raw(Hz)"] = onset_freq
        out.at[index, "红点频率_used(Hz)"] = onset_freq
        out.at[index, "时间误差(ms)"] = official_time_error_ms
        out.at[index, "音准误差(cents)"] = pitch_error
        out.at[index, "音准等级"] = classify_pitch_level(pitch_error)
        out.at[index, "节奏等级"] = (
            classify_rhythm_level(official_time_error_ms)
            if np.isfinite(official_time_error_ms)
            else ""
        )
        out.at[index, "后删标记"] = ""
        out.at[index, "recovery_onset_id"] = onset_id
        out.at[index, "recovery_cost"] = cost
        out.at[index, "recovery_time_error_ms"] = recovery_time_error_ms
        out.at[index, "recovery_pitch_error_cents"] = pitch_error
        out.at[index, "recovery_result"] = "recovered"
        used_ids.add(onset_id)

    return out


def recover_structured_weak_candidates(df_align, df_sc, df_on, df_weak):
    """Recover two evidence-backed gaps without admitting weak rows globally.

    1) voiced_reentry_weak: requires two reliable anchors and a unique,
       score-compatible candidate inside the normal local prediction window.
    2) gradual_platform_transition_weak: requires an immediately preceding
       matched grace at the same score time, with left/right platform pitch
       evidence agreeing with the grace/main score pitches.
    """
    out = df_align.copy()
    by_frame = {}
    allowed_sources = {VOICED_REENTRY_WEAK_SOURCE, GRADUAL_PLATFORM_WEAK_SOURCE}
    for _, onset in df_weak.iterrows():
        source = str(onset.get("pre_src", "")).strip().lower()
        level = str(onset.get("candidate_level", "")).strip().lower()
        onset_id = str(onset.get(ONSET_ID_COL, "")).strip()
        frame = safe_float(onset.get("note_frame"))
        onset_time = safe_float(onset.get(ONSET_T_COL))
        onset_freq = safe_float(onset.get(ONSET_F0_COL))
        confidence = safe_float(onset.get("candidate_confidence"))
        if (
            source not in allowed_sources
            or level != "weak"
            or not onset_id
            or not np.isfinite(frame)
            or frame < 0
            or not np.isfinite(onset_time)
            or not np.isfinite(onset_freq)
            or onset_freq <= 0
            or not np.isfinite(confidence)
            or confidence < STRUCTURED_WEAK_MIN_CONFIDENCE
        ):
            continue
        key = int(frame)
        item = (confidence, onset_id, onset_time, onset_freq, source, onset)
        existing = by_frame.get(key)
        if existing is None or confidence > existing[0]:
            by_frame[key] = item
    normal_by_frame = {}
    for _, onset in df_on.iterrows():
        onset_id = str(onset.get(ONSET_ID_COL, "")).strip()
        frame = safe_float(onset.get("note_frame"))
        onset_time = safe_float(onset.get(ONSET_T_COL))
        onset_freq = safe_float(onset.get(ONSET_F0_COL))
        source = str(onset.get("pre_src", "")).strip().lower()
        if (
            not onset_id
            or "black" in {part.strip() for part in source.split("+")}
            or not np.isfinite(frame)
            or frame < 0
            or not np.isfinite(onset_time)
            or not np.isfinite(onset_freq)
            or onset_freq <= 0
        ):
            continue
        key = int(frame)
        item = (1.0, onset_id, onset_time, onset_freq, EXISTING_STRONG_GRACE_SOURCE, onset)
        existing = normal_by_frame.get(key)
        if existing is None or (onset_time, onset_id) < (existing[2], existing[1]):
            normal_by_frame[key] = item

    used_ids = {
        str(row.get(ONSET_ID_COL, "")).strip()
        for _, row in out.iterrows()
        if _valid_matched_row(row) and str(row.get(ONSET_ID_COL, "")).strip()
    }

    for index in range(len(out)):
        if _recovery_status_priority(out.at[index, "对齐状态"]) not in {1, 2}:
            continue
        score_row = df_sc.loc[index]
        align_flag = str(score_row.get("参与起音对齐", "是")).strip().lower()
        score_freq = safe_float(out.at[index, "谱面频率(Hz)"])
        if align_flag in {"否", "no", "false", "0"} or not np.isfinite(score_freq) or score_freq <= 0:
            continue

        left, right = _nearest_recovery_anchors(out, index)
        if left is None:
            continue
        left_time = safe_float(out.at[left, "红点时间(s)"])
        right_time = safe_float(out.at[right, "红点时间(s)"]) if right is not None else np.nan
        original_status = str(out.at[index, "对齐状态"])
        qualified = []

        # Strict bilateral voiced re-entry recovery.
        if right is not None and np.isfinite(right_time) and right_time > left_time:
            predicted, window_ms, prediction_source = _predict_recovery_time(out, index, left, right)
            if np.isfinite(predicted):
                window_ms = min(250.0, max(80.0, float(window_ms)))
                for confidence, onset_id, onset_time, onset_freq, source, onset_row in by_frame.values():
                    if source != VOICED_REENTRY_WEAK_SOURCE or onset_id in used_ids:
                        continue
                    if not (left_time < onset_time < right_time):
                        continue
                    time_error_ms = (onset_time - predicted) * 1000.0
                    if abs(time_error_ms) > window_ms:
                        continue
                    pitch_error = cents_error(score_freq, onset_freq)
                    if abs(pitch_error) > STRUCTURED_WEAK_MAX_PITCH_CENTS:
                        continue
                    cost = (
                        0.55 * abs(time_error_ms) / window_ms
                        + 0.35 * abs(pitch_error) / STRUCTURED_WEAK_MAX_PITCH_CENTS
                        + 0.10 * (1.0 - confidence)
                    )
                    if cost <= STRUCTURED_WEAK_MAX_COST:
                        qualified.append((
                            cost, onset_id, onset_time, onset_freq, confidence,
                            pitch_error, time_error_ms, onset_row,
                            "dp3_voiced_reentry_weak", prediction_source,
                            window_ms,
                        ))

        # Matched grace -> unmatched main at the same score time.  The weak
        # candidate must acoustically connect the grace platform to the main
        # platform; time alone is never sufficient.
        previous = index - 1
        if previous >= 0 and _valid_matched_row(out.loc[previous]):
            previous_score_time = safe_float(out.at[previous, "谱面时间_原始(s)"])
            current_score_time = safe_float(out.at[index, "谱面时间_原始(s)"])
            previous_grace = str(out.at[previous, "装饰音判定"]).strip() == "装饰音"
            current_main = str(out.at[index, "装饰音判定"]).strip() == "主音"
            same_score_time = (
                np.isfinite(previous_score_time)
                and np.isfinite(current_score_time)
                and abs(previous_score_time - current_score_time) <= 1e-9
            )
            if previous_grace and current_main and same_score_time and left == previous:
                previous_score_freq = safe_float(out.at[previous, "谱面频率(Hz)"])
                grace_candidates = list(by_frame.values()) + list(normal_by_frame.values())
                for confidence, onset_id, onset_time, onset_freq, source, onset_row in grace_candidates:
                    if source not in {GRADUAL_PLATFORM_WEAK_SOURCE, EXISTING_STRONG_GRACE_SOURCE} or onset_id in used_ids:
                        continue
                    if onset_time <= left_time or (right is not None and not onset_time < right_time):
                        continue
                    pitch_error = cents_error(score_freq, onset_freq)
                    if source == GRADUAL_PLATFORM_WEAK_SOURCE:
                        platform_left = safe_float(onset_row.get("platform_left_f0"))
                        platform_right = safe_float(onset_row.get("platform_right_f0"))
                        if not (
                            np.isfinite(previous_score_freq)
                            and np.isfinite(platform_left)
                            and platform_left > 0
                            and np.isfinite(platform_right)
                            and platform_right > 0
                        ):
                            continue
                        left_pitch_error = cents_error(previous_score_freq, platform_left)
                        right_pitch_error = cents_error(score_freq, platform_right)
                        if max(
                            abs(left_pitch_error), abs(right_pitch_error), abs(pitch_error)
                        ) > STRUCTURED_WEAK_MAX_PITCH_CENTS:
                            continue
                    else:
                        left_pitch_error = 0.0
                        right_pitch_error = pitch_error
                        if abs(pitch_error) > STRUCTURED_WEAK_MAX_PITCH_CENTS:
                            continue
                    span_ms = (
                        max(1.0, (right_time - left_time) * 1000.0)
                        if right is not None and np.isfinite(right_time)
                        else max(1000.0, (onset_time - left_time) * 1000.0)
                    )
                    relative_time = min(1.0, max(0.0, (onset_time - left_time) * 1000.0 / span_ms))
                    if source == GRADUAL_PLATFORM_WEAK_SOURCE:
                        cost = (
                            0.30 * abs(left_pitch_error) / STRUCTURED_WEAK_MAX_PITCH_CENTS
                            + 0.35 * abs(right_pitch_error) / STRUCTURED_WEAK_MAX_PITCH_CENTS
                            + 0.20 * abs(pitch_error) / STRUCTURED_WEAK_MAX_PITCH_CENTS
                            + 0.10 * (1.0 - confidence)
                            + 0.05 * relative_time
                        )
                        prediction_source = "grace_main_platform_identity"
                    else:
                        cost = (
                            0.90 * abs(pitch_error) / STRUCTURED_WEAK_MAX_PITCH_CENTS
                            + 0.10 * relative_time
                        )
                        prediction_source = "grace_main_existing_strong_candidate"
                    if cost <= STRUCTURED_WEAK_MAX_COST:
                        qualified.append((
                            cost, onset_id, onset_time, onset_freq, confidence,
                            pitch_error, (onset_time - left_time) * 1000.0,
                            onset_row, "dp3_grace_main_plateau",
                            prediction_source, span_ms,
                        ))

        qualified.sort(key=lambda item: (item[0], item[2], item[1]))
        if not qualified:
            continue
        if (
            len(qualified) > 1
            and qualified[1][0] - qualified[0][0] < STRUCTURED_WEAK_AMBIGUITY_MARGIN
        ):
            continue

        (
            cost, onset_id, onset_time, onset_freq, confidence, pitch_error,
            recovery_time_error_ms, onset_row, stage, prediction_source,
            search_window_ms,
        ) = qualified[0]
        shifted_score_time = safe_float(out.at[index, "谱面时间_平移后(s)"])
        official_time_error_ms = (
            (onset_time - shifted_score_time) * 1000.0
            if np.isfinite(shifted_score_time)
            else np.nan
        )
        out.at[index, "对齐状态"] = "对齐成功"
        out.at[index, "alignment_status"] = "对齐成功"
        out.at[index, ONSET_ID_COL] = onset_id
        out.at[index, ORIGINAL_ONSET_INDEX_COL] = onset_row.get(ORIGINAL_ONSET_INDEX_COL, np.nan)
        out.at[index, CURRENT_ONSET_INDEX_COL] = onset_row.get(CURRENT_ONSET_INDEX_COL, np.nan)
        out.at[index, "红点时间(s)"] = onset_time
        out.at[index, "红点频率_raw(Hz)"] = onset_freq
        out.at[index, "红点频率_used(Hz)"] = onset_freq
        out.at[index, "时间误差(ms)"] = official_time_error_ms
        out.at[index, "音准误差(cents)"] = pitch_error
        out.at[index, "音准等级"] = classify_pitch_level(pitch_error)
        out.at[index, "节奏等级"] = classify_rhythm_level(official_time_error_ms)
        out.at[index, "后删标记"] = ""
        out.at[index, "recovery_stage"] = stage
        out.at[index, "recovery_original_status"] = original_status
        out.at[index, "recovery_onset_id"] = onset_id
        out.at[index, "recovery_cost"] = cost
        out.at[index, "recovery_time_error_ms"] = recovery_time_error_ms
        out.at[index, "recovery_pitch_error_cents"] = pitch_error
        out.at[index, "recovery_predicted_time"] = onset_time
        out.at[index, "recovery_search_window_ms"] = search_window_ms
        out.at[index, "recovery_candidate_count"] = len(qualified)
        out.at[index, "recovery_prediction_source"] = prediction_source
        out.at[index, "recovery_result"] = f"recovered_confidence_{confidence:.3f}"
        used_ids.add(onset_id)
        print(
            f"[DP3 structured weak] score[{index}] uses {onset_id} "
            f"at {onset_time:.3f}s via {stage}"
        )

    return out


def recover_platform_weak_bridge(df_align, df_sc, df_weak):
    """Close one direct recovery-to-gap bridge with a high-confidence weak red.

    This is intentionally *not* another DP pass.  DP1/DP2 and the main DP3
    pass only see strong candidates.  Here a weak platform candidate is usable
    only when the immediately preceding score note was already conservatively
    recovered, the current row is still unmatched, and timing/pitch/order all
    agree with the local anchors.  That keeps a weak candidate from rerouting
    an otherwise normal phrase.
    """
    out = df_align.copy()
    if not ENABLE_PLATFORM_WEAK_GAP_RESCUE or df_weak is None or df_weak.empty:
        return out

    weak_by_frame = {}
    for _, onset in df_weak.iterrows():
        source = str(onset.get("pre_src", "")).strip().lower()
        level = str(onset.get("candidate_level", "")).strip().lower()
        onset_id = str(onset.get(ONSET_ID_COL, "")).strip()
        frame = safe_float(onset.get("note_frame"))
        onset_time = safe_float(onset.get(ONSET_T_COL))
        onset_freq = safe_float(onset.get(ONSET_F0_COL))
        confidence = safe_float(onset.get("candidate_confidence"))
        if (
            source != "platform_transition_weak"
            or level != "weak"
            or not onset_id
            or not np.isfinite(frame)
            or frame < 0
            or not np.isfinite(onset_time)
            or not np.isfinite(onset_freq)
            or onset_freq <= 0
            or not np.isfinite(confidence)
            or confidence < PLATFORM_WEAK_MIN_CONFIDENCE
        ):
            continue
        key = int(frame)
        existing = weak_by_frame.get(key)
        if existing is None or confidence > existing[0]:
            weak_by_frame[key] = (confidence, onset_id, onset_time, onset_freq, onset)

    if not weak_by_frame:
        return out

    used_ids = {
        str(row.get(ONSET_ID_COL, "")).strip()
        for _, row in out.iterrows()
        if _valid_matched_row(row) and str(row.get(ONSET_ID_COL, "")).strip()
    }

    for index in range(1, len(out)):
        # A direct bridge is deliberately narrower than generic recovery:
        # recovered previous note + immediately following still-unmatched note.
        # Both a post-drop unmatched row (priority 1) and a plain unmatched
        # row (priority 2) are eligible; conflict rows are deliberately not.
        if _recovery_status_priority(out.at[index, "对齐状态"]) not in {1, 2}:
            continue
        previous = index - 1
        if not _is_conservative_recovery_row(out.loc[previous]):
            continue
        score_row = df_sc.loc[index]
        align_flag = str(score_row.get("参与起音对齐", "是")).strip().lower()
        score_freq = safe_float(out.at[index, "谱面频率(Hz)"])
        if align_flag in {"否", "no", "false", "0"} or not np.isfinite(score_freq) or score_freq <= 0:
            continue

        left, right = _nearest_recovery_anchors(out, index)
        if left != previous:
            continue
        predicted, window_ms, prediction_source = _predict_recovery_time(out, index, left, right)
        if not np.isfinite(predicted):
            continue
        window_ms = min(float(PLATFORM_WEAK_MAX_TIME_WINDOW_MS), max(80.0, float(window_ms)))
        left_time = safe_float(out.at[left, "红点时间(s)"])
        right_time = safe_float(out.at[right, "红点时间(s)"]) if right is not None else np.nan

        qualified = []
        for confidence, onset_id, onset_time, onset_freq, onset_row in weak_by_frame.values():
            if onset_id in used_ids or not (onset_time > left_time):
                continue
            if right is not None and not (onset_time < right_time):
                continue
            time_error_ms = (onset_time - predicted) * 1000.0
            if abs(time_error_ms) > window_ms:
                continue
            pitch_error = cents_error(score_freq, onset_freq)
            if abs(pitch_error) > PLATFORM_WEAK_MAX_PITCH_CENTS:
                continue
            cost = (
                0.55 * abs(time_error_ms) / window_ms
                + 0.35 * abs(pitch_error) / PLATFORM_WEAK_MAX_PITCH_CENTS
                + 0.10 * (1.0 - confidence)
            )
            qualified.append((cost, onset_id, onset_time, onset_freq, confidence, pitch_error, time_error_ms, onset_row))

        qualified.sort(key=lambda item: (item[0], item[2], item[1]))
        if not qualified:
            continue
        if (
            len(qualified) > 1
            and qualified[1][0] - qualified[0][0] < PLATFORM_WEAK_AMBIGUITY_MARGIN
        ):
            continue

        cost, onset_id, onset_time, onset_freq, confidence, pitch_error, recovery_time_error_ms, onset_row = qualified[0]
        shifted_score_time = safe_float(out.at[index, "谱面时间_平移后(s)"])
        official_time_error_ms = (
            (onset_time - shifted_score_time) * 1000.0
            if np.isfinite(shifted_score_time)
            else np.nan
        )
        original_status = str(out.at[index, "对齐状态"])
        out.at[index, "对齐状态"] = "对齐成功"
        out.at[index, "alignment_status"] = "对齐成功"
        out.at[index, ONSET_ID_COL] = onset_id
        out.at[index, ORIGINAL_ONSET_INDEX_COL] = onset_row.get(ORIGINAL_ONSET_INDEX_COL, np.nan)
        out.at[index, CURRENT_ONSET_INDEX_COL] = onset_row.get(CURRENT_ONSET_INDEX_COL, np.nan)
        out.at[index, "红点时间(s)"] = onset_time
        out.at[index, "红点频率_raw(Hz)"] = onset_freq
        out.at[index, "红点频率_used(Hz)"] = onset_freq
        out.at[index, "时间误差(ms)"] = official_time_error_ms
        out.at[index, "音准误差(cents)"] = pitch_error
        out.at[index, "音准等级"] = classify_pitch_level(pitch_error)
        out.at[index, "节奏等级"] = (
            classify_rhythm_level(official_time_error_ms)
            if np.isfinite(official_time_error_ms)
            else ""
        )
        out.at[index, "后删标记"] = ""
        out.at[index, "recovery_stage"] = "dp3_platform_weak_bridge"
        out.at[index, "recovery_original_status"] = original_status
        out.at[index, "recovery_onset_id"] = onset_id
        out.at[index, "recovery_cost"] = cost
        out.at[index, "recovery_time_error_ms"] = recovery_time_error_ms
        out.at[index, "recovery_pitch_error_cents"] = pitch_error
        out.at[index, "recovery_predicted_time"] = predicted
        out.at[index, "recovery_search_window_ms"] = window_ms
        out.at[index, "recovery_candidate_count"] = len(qualified)
        out.at[index, "recovery_prediction_source"] = prediction_source
        out.at[index, "recovery_result"] = f"weak_bridge_confidence_{confidence:.3f}"
        used_ids.add(onset_id)

        # The preceding red is no longer an isolated fallback: together the
        # two ordered red events form a normal local pair.  Preserve recovery
        # provenance in its audit columns while exposing normal final status.
        old_stage = str(out.at[previous, "recovery_stage"]).strip()
        out.at[previous, "对齐状态"] = "对齐成功"
        out.at[previous, "alignment_status"] = "对齐成功"
        out.at[previous, "recovery_stage"] = (
            f"{old_stage}|normalized_by_platform_weak_bridge".strip("|")
        )
        out.at[previous, "recovery_result"] = "normalized_by_platform_weak_bridge"
        print(
            f"[DP3 weak bridge] score[{previous}] + score[{index}] "
            f"uses {onset_id} at {onset_time:.3f}s (confidence={confidence:.3f})"
        )

    return out


def _score_measure_key(value):
    """Extract the MusicXML measure identity without depending on a title."""
    for token in str(value).strip().split("_"):
        if len(token) > 1 and token[0] == "M" and token[1:].isdigit():
            return token
    return ""


def anchor_terminal_suffix_to_observed_black(
    df_align,
    df_sc,
    df_on,
    observed_black_time=np.nan,
):
    """Re-open an unresolved final-measure suffix against the real end marker.

    This is an endpoint-boundary correction, not a piece-specific repair.  It
    is enabled only when the final MusicXML measure still contains an explicit
    unresolved score event.  All following main notes are re-assigned jointly
    in score/time order; grace notes may remain unmatched.  The objective uses
    the existing pitch gate and interval-ratio cost, with the observed black
    marker as the terminal anchor.
    """
    out = df_align.copy()
    if not np.isfinite(observed_black_time) or len(out) == 0:
        return out
    if "score_note_id" not in df_sc.columns:
        return out

    final_measure = _score_measure_key(df_sc.iloc[-1].get("score_note_id", ""))
    if not final_measure:
        return out
    unresolved = [
        index
        for index in range(len(out))
        if not _valid_matched_row(out.loc[index])
        and _score_measure_key(df_sc.loc[index].get("score_note_id", "")) == final_measure
    ]
    if not unresolved:
        return out
    start = min(unresolved)

    left = None
    for index in range(start - 1, -1, -1):
        if _valid_matched_row(out.loc[index]):
            left = index
            break
    if left is None:
        return out
    left_time = safe_float(out.at[left, "红点时间(s)"])
    left_id = str(out.at[left, ONSET_ID_COL]).strip()
    left_on_idx = onset_index_by_id(df_on, left_id)
    if (
        left_on_idx is None
        or not np.isfinite(left_time)
        or not (observed_black_time > left_time)
    ):
        return out

    score_t = get_score_time_array(df_sc)
    score_dur = get_score_duration_array(df_sc)
    onset_t = get_onset_time_array(df_on)
    sc_f = safe_numeric(df_sc[SCORE_F0_COL]).to_numpy(dtype=float)
    sc_m = midi_round(sc_f)
    on_f = safe_numeric(df_on[ONSET_F0_COL]).to_numpy(dtype=float)
    if score_t is None or score_dur is None or onset_t is None:
        return out
    score_black_time = float(score_t[-1] + score_dur[-1])
    if not np.isfinite(score_black_time):
        return out

    suffix_indices = set(range(start, len(out)))
    blocked_ids = {
        str(row.get(ONSET_ID_COL, "")).strip()
        for index, row in out.iterrows()
        if index not in suffix_indices
        and _valid_matched_row(row)
        and str(row.get(ONSET_ID_COL, "")).strip()
    }
    candidates = []
    for on_idx, onset in df_on.iterrows():
        onset_id = str(onset.get(ONSET_ID_COL, "")).strip()
        onset_time = safe_float(onset.get(ONSET_T_COL))
        onset_freq = safe_float(onset.get(ONSET_F0_COL))
        source = str(onset.get("pre_src", "")).strip().lower()
        if (
            not onset_id
            or onset_id in blocked_ids
            or "black" in {part.strip() for part in source.split("+")}
            or not (np.isfinite(onset_time) and left_time < onset_time < observed_black_time)
            or not (np.isfinite(onset_freq) and onset_freq > 0)
        ):
            continue
        candidates.append((int(on_idx), onset_id, onset_time, onset_freq, onset))
    candidates.sort(key=lambda item: (item[2], item[1]))
    if not candidates:
        return out

    # State key: (last matched score index, candidate position, match count).
    # Values keep total interval cost and the assignment path for the suffix.
    states = {(left, -1, 0): (0.0, tuple())}
    for score_index in range(start, len(out)):
        align_flag = str(df_sc.loc[score_index].get("参与起音对齐", "是")).strip().lower()
        optional = is_grace_note(score_index, df_sc) or align_flag in {"否", "no", "false", "0"}
        next_states = {}
        for (last_score_index, last_position, match_count), (total_cost, path) in states.items():
            if optional:
                key = (last_score_index, last_position, match_count)
                value = (total_cost, path + (None,))
                if key not in next_states or value[0] < next_states[key][0]:
                    next_states[key] = value

            last_on_index = left_on_idx if last_position < 0 else candidates[last_position][0]
            last_on_time = left_time if last_position < 0 else candidates[last_position][2]
            score_gap = local_transition_score_gap(
                last_score_index, score_index, score_t, df_sc
            )
            if not np.isfinite(score_gap) or score_gap <= 0:
                continue

            for position in range(last_position + 1, len(candidates)):
                on_idx, _onset_id, onset_time, _onset_freq, _onset_row = candidates[position]
                if not candidate_pitch_ok(score_index, on_idx, sc_f, on_f):
                    continue
                if not main_same_note_gap_ok(
                    last_score_index,
                    score_index,
                    last_on_index,
                    on_idx,
                    df_sc,
                    score_t,
                    onset_t,
                    sc_m,
                ):
                    continue
                edge_cost = interval_ratio_error(onset_time - last_on_time, score_gap)
                if not np.isfinite(edge_cost):
                    continue
                key = (score_index, position, match_count + 1)
                value = (total_cost + edge_cost, path + (on_idx,))
                old = next_states.get(key)
                if old is None or (value[0], value[1]) < (old[0], old[1]):
                    next_states[key] = value
        states = next_states
        if not states:
            return out

    scored = []
    for (last_score_index, last_position, match_count), (total_cost, path) in states.items():
        if last_position < 0 or match_count <= 0:
            continue
        last_on_time = candidates[last_position][2]
        score_gap = score_black_time - score_t[last_score_index]
        end_cost = interval_ratio_error(observed_black_time - last_on_time, score_gap)
        if not np.isfinite(end_cost):
            continue
        mean_cost = (total_cost + end_cost) / (match_count + 1)
        if mean_cost <= REPEATED_PITCH_CONFLICT_MAX_COST:
            scored.append((mean_cost, total_cost + end_cost, path))
    if not scored:
        print(
            f"[DP3 terminal anchor] final measure {final_measure}: "
            "no complete main-note path; keep unmatched state"
        )
        return out

    scored.sort(
        key=lambda item: (
            item[0],
            item[1],
            tuple(-1 if value is None else int(value) for value in item[2]),
        )
    )
    best_cost, _best_total, best_path = scored[0]

    recovery_defaults = {
        "alignment_status": out.get("对齐状态", pd.Series("", index=out.index)),
        "recovery_stage": "",
        "recovery_original_status": "",
        "recovery_onset_id": "",
        "recovery_cost": np.nan,
        "recovery_time_error_ms": np.nan,
        "recovery_pitch_error_cents": np.nan,
        "recovery_predicted_time": np.nan,
        "recovery_search_window_ms": np.nan,
        "recovery_candidate_count": 0,
        "recovery_prediction_source": "",
        "recovery_result": "",
    }
    for column, default in recovery_defaults.items():
        if column not in out.columns:
            out[column] = default

    onset_row_by_index = {int(index): row for index, row in df_on.iterrows()}
    old_assignments = []
    new_assignments = []
    for score_index, on_idx in zip(range(start, len(out)), best_path):
        original_status = str(out.at[score_index, "对齐状态"])
        old_assignments.append(str(out.at[score_index, ONSET_ID_COL]).strip())
        out.at[score_index, "recovery_stage"] = "dp3_observed_black_terminal_anchor"
        out.at[score_index, "recovery_original_status"] = original_status
        out.at[score_index, "recovery_cost"] = best_cost
        out.at[score_index, "recovery_search_window_ms"] = (
            observed_black_time - left_time
        ) * 1000.0
        out.at[score_index, "recovery_candidate_count"] = len(candidates)
        out.at[score_index, "recovery_prediction_source"] = (
            "final_musicxml_measure_to_observed_black"
        )

        if on_idx is None:
            new_assignments.append("")
            out.at[score_index, "对齐状态"] = "unmatched"
            out.at[score_index, "alignment_status"] = "unmatched"
            out.at[score_index, ONSET_ID_COL] = ""
            out.at[score_index, ORIGINAL_ONSET_INDEX_COL] = np.nan
            out.at[score_index, CURRENT_ONSET_INDEX_COL] = np.nan
            out.at[score_index, "红点时间(s)"] = np.nan
            out.at[score_index, "红点频率_raw(Hz)"] = np.nan
            out.at[score_index, "红点频率_used(Hz)"] = np.nan
            out.at[score_index, "时间误差(ms)"] = np.nan
            out.at[score_index, "音准误差(cents)"] = np.nan
            out.at[score_index, "音准等级"] = ""
            out.at[score_index, "节奏等级"] = ""
            out.at[score_index, "后删标记"] = ""
            out.at[score_index, "recovery_onset_id"] = ""
            out.at[score_index, "recovery_pitch_error_cents"] = np.nan
            out.at[score_index, "recovery_result"] = "optional_grace_unmatched"
            continue

        onset = onset_row_by_index[int(on_idx)]
        onset_id = str(onset.get(ONSET_ID_COL, "")).strip()
        onset_time = safe_float(onset.get(ONSET_T_COL))
        onset_freq = safe_float(onset.get(ONSET_F0_COL))
        pitch_error = cents_error(sc_f[score_index], onset_freq)
        shifted_score_time = safe_float(out.at[score_index, "谱面时间_平移后(s)"])
        official_time_error_ms = (
            (onset_time - shifted_score_time) * 1000.0
            if np.isfinite(shifted_score_time)
            else np.nan
        )
        new_assignments.append(onset_id)
        out.at[score_index, "对齐状态"] = "对齐成功"
        out.at[score_index, "alignment_status"] = "对齐成功"
        out.at[score_index, ONSET_ID_COL] = onset_id
        out.at[score_index, ORIGINAL_ONSET_INDEX_COL] = onset.get(
            ORIGINAL_ONSET_INDEX_COL, np.nan
        )
        out.at[score_index, CURRENT_ONSET_INDEX_COL] = onset.get(
            CURRENT_ONSET_INDEX_COL, np.nan
        )
        out.at[score_index, "红点时间(s)"] = onset_time
        out.at[score_index, "红点频率_raw(Hz)"] = onset_freq
        out.at[score_index, "红点频率_used(Hz)"] = onset_freq
        out.at[score_index, "时间误差(ms)"] = official_time_error_ms
        out.at[score_index, "音准误差(cents)"] = pitch_error
        out.at[score_index, "音准等级"] = classify_pitch_level(pitch_error)
        out.at[score_index, "节奏等级"] = classify_rhythm_level(
            official_time_error_ms
        )
        out.at[score_index, "后删标记"] = ""
        out.at[score_index, "recovery_onset_id"] = onset_id
        out.at[score_index, "recovery_pitch_error_cents"] = pitch_error
        out.at[score_index, "recovery_result"] = "terminal_joint_reassigned"

    print(
        f"[DP3 terminal anchor] final measure {final_measure}, score[{start}:{len(out)-1}], "
        f"black={observed_black_time:.3f}s, cost={best_cost:.4f}, "
        f"onset {old_assignments} -> {new_assignments}"
    )
    return out


def append_duration_ratio_columns(df_align, observed_black_time=np.nan):
    if len(df_align) == 0:
        return df_align

    actual_gaps = []
    score_gaps = []
    ratios = []
    ratio_pcts = []
    ratio_levels = []

    # 时值统计必须使用“平移后的谱面时间”。
    # 否则最后一个虚拟 score black 会停在原始谱面终点，导致尾音实际推进时长变负。
    if "谱面时间_平移后(s)" in df_align.columns:
        score_times = pd.to_numeric(df_align["谱面时间_平移后(s)"], errors="coerce").to_numpy(dtype=float)
    else:
        score_times = pd.to_numeric(df_align["谱面时间_原始(s)"], errors="coerce").to_numpy(dtype=float)
    onset_times = pd.to_numeric(df_align["红点时间(s)"], errors="coerce").to_numpy(dtype=float)
    score_durs = pd.to_numeric(df_align["持续时间(s)"], errors="coerce").to_numpy(dtype=float)

    if "装饰音判定" in df_align.columns:
        grace_flags = df_align["装饰音判定"].fillna("").astype(str).to_list()
    else:
        grace_flags = [""] * len(df_align)

    if "BPM" in df_align.columns:
        bpm_vals = pd.to_numeric(df_align["BPM"], errors="coerce").to_numpy(dtype=float)
    else:
        bpm_vals = np.full(len(df_align), np.nan, dtype=float)

    score_black_time = np.nan
    if len(df_align) > 0 and np.isfinite(score_times[-1]) and np.isfinite(score_durs[-1]):
        score_black_time = score_times[-1] + score_durs[-1]
    actual_black_time = (
        float(observed_black_time)
        if np.isfinite(observed_black_time)
        else score_black_time
    )

    for i in range(len(df_align)):
        current_valid = _valid_matched_row(df_align.loc[i]) and np.isfinite(onset_times[i])
        next_valid_index = None
        if current_valid:
            for candidate in range(i + 1, len(df_align)):
                if (
                    _valid_matched_row(df_align.loc[candidate])
                    and np.isfinite(onset_times[candidate])
                ):
                    next_valid_index = candidate
                    break

        if not current_valid:
            actual_gap = np.nan
        elif next_valid_index is not None:
            actual_gap = onset_times[next_valid_index] - onset_times[i]
        else:
            actual_gap = (
                actual_black_time - onset_times[i]
                if np.isfinite(actual_black_time)
                else np.nan
            )

        is_grace = (grace_flags[i].strip() == "装饰音")
        if is_grace:
            bpm_val = bpm_vals[i]
            score_gap = 6.0 / bpm_val if np.isfinite(bpm_val) and bpm_val > 0 else np.nan
        else:
            if not current_valid:
                score_gap = np.nan
            elif next_valid_index is not None:
                score_gap = score_times[next_valid_index] - score_times[i]
            else:
                score_gap = (
                    score_black_time - score_times[i]
                    if np.isfinite(score_black_time)
                    else np.nan
                )

        if not np.isfinite(actual_gap) or not np.isfinite(score_gap) or score_gap <= 0:
            ratio = np.nan
            ratio_pct = np.nan
            ratio_level = "F"
        else:
            ratio = actual_gap / score_gap
            ratio_pct = (ratio - 1.0) * 100.0
            ratio_level = classify_duration_ratio_level(ratio, is_grace)

        actual_gaps.append(actual_gap)
        score_gaps.append(score_gap)
        ratios.append(ratio)
        ratio_pcts.append(ratio_pct)
        ratio_levels.append(ratio_level)

    df_align["实际推进时长(s)"] = actual_gaps
    df_align["理论推进时长(s)"] = score_gaps
    df_align["时值比"] = ratios
    df_align["时值偏差率(%)"] = ratio_pcts
    df_align["时值等级"] = ratio_levels
    return df_align


def build_metrics_table(df_align):
    total = len(df_align)
    valid_mask = (df_align["对齐状态"] == "对齐成功")
    df_valid = df_align[valid_mask].copy()

    success = len(df_valid)
    dropped = int((df_align["对齐状态"] == "未匹配(后删)").sum())
    success_rate = success / total * 100.0 if total > 0 else np.nan

    time_err = safe_numeric(df_valid["时间误差(ms)"]) if success > 0 else pd.Series(dtype=float)
    pitch_err = safe_numeric(df_valid["音准误差(cents)"]) if success > 0 else pd.Series(dtype=float)
    pitch_counts = df_valid["音准等级"].value_counts() if success > 0 else pd.Series(dtype=int)
    rhythm_counts = df_valid["节奏等级"].value_counts() if success > 0 else pd.Series(dtype=int)

    score_shift_val = round(float(df_align["score_shift(s)"].iloc[0]) if total > 0 else 0.0, 6)

    return pd.DataFrame([{
        "总谱面音数": total,
        "后删未匹配数": dropped,
        "最终匹配成功数": success,
        "最终匹配率(%)": round(success_rate, 2),
        "score_shift(s)": score_shift_val,
        "平均时间误差(ms)": round(time_err.mean(), 2) if success > 0 else np.nan,
        "时间误差中位数(ms)": round(time_err.median(), 2) if success > 0 else np.nan,
        "时间误差标准差(ms)": round(time_err.std(), 2) if success > 0 else np.nan,
        "平均绝对时间误差(ms)": round(time_err.abs().mean(), 2) if success > 0 else np.nan,
        "平均音准误差(cents)": round(pitch_err.mean(), 2) if success > 0 else np.nan,
        "音准误差中位数(cents)": round(pitch_err.median(), 2) if success > 0 else np.nan,
        "音准误差标准差(cents)": round(pitch_err.std(), 2) if success > 0 else np.nan,
        "平均绝对音准误差(cents)": round(pitch_err.abs().mean(), 2) if success > 0 else np.nan,
        "音准S(<=8c)个数": int(pitch_counts.get("S", 0)),
        "音准A(<=20c)个数": int(pitch_counts.get("A", 0)),
        "音准B(<=35c)个数": int(pitch_counts.get("B", 0)),
        "音准C(<=50c)个数": int(pitch_counts.get("C", 0)),
        "音准D(<=75c)个数": int(pitch_counts.get("D", 0)),
        "音准E(<=100c)个数": int(pitch_counts.get("E", 0)),
        "音准F(后删)个数": dropped,
        "节奏优秀(<=80ms)个数": int(rhythm_counts.get("优秀", 0)),
        "节奏可接受(<=150ms)个数": int(rhythm_counts.get("可接受", 0)),
        "节奏偏差较大(>150ms)个数": int(rhythm_counts.get("偏差较大", 0)),
        "低音八度修正个数": 0,
        "高八度等效匹配个数": 0,
        "高八度候选但未放开个数": 0,
    }])


# ============================================================
# 主流程
# ============================================================
def run_dp3(pairs, df_sc, df_on, sc_m, sc_f, observed_black_time=np.nan):
    pairs = ornament_back_refine(pairs, df_sc, df_on, sc_m, sc_f)
    pairs = remove_duplicate_onset_pairs(pairs, len(df_on))
    pairs = [(si, oj, "否", False) for si, oj, _, _ in pairs if 0 <= oj < len(df_on)]

    print("\n[阶段3] 同音竞争重排...")
    pairs, locked_score_indices = same_note_group_reassign(
        pairs,
        df_sc,
        df_on,
        sc_m,
        sc_f,
        return_locked=True,
        observed_black_time=observed_black_time,
    )
    pairs = remove_duplicate_onset_pairs(pairs, len(df_on))
    pairs = [(si, oj, "否", False) for si, oj, _, _ in pairs if 0 <= oj < len(df_on)]

    print("\n[阶段4] 从前往后前卷...")
    # First preserve the exact legacy result.  The protected pass is accepted
    # only for score entries whose completed outcome improves over it.
    baseline_forward_pairs = forward_early_refine(pairs, df_sc, df_on, sc_f)
    if locked_score_indices:
        protected_forward_pairs = forward_early_refine(
            pairs, df_sc, df_on, sc_f, locked_score_indices=locked_score_indices
        )
        accepted_locks = accept_same_note_forward_locks(
            baseline_forward_pairs,
            protected_forward_pairs,
            locked_score_indices,
            df_sc,
            df_on,
            sc_f,
        )
        if accepted_locks:
            print(f"  accepted same-note forward locks: {len(accepted_locks)}")
            pairs = forward_early_refine(
                pairs, df_sc, df_on, sc_f, locked_score_indices=accepted_locks
            )
        else:
            print("  no same-note forward locks passed the completed-output check")
            pairs = baseline_forward_pairs
    else:
        pairs = baseline_forward_pairs
    pairs = remove_duplicate_onset_pairs(pairs, len(df_on))
    pairs = [(si, oj, "否", False) for si, oj, _, _ in pairs if 0 <= oj < len(df_on)]
    return pairs


def run_output_stage(
    df_sc,
    df_on,
    pairs,
    df_weak=None,
    observed_black_time=np.nan,
):
    print("\n[阶段5] 生成结果表与统计...")
    df_align = build_align_table(df_sc, df_on, pairs)
    df_align = complete_alignment_rows(df_align, df_sc)
    df_align = reassign_repeated_pitch_conflicts(df_align, df_sc, df_on)
    df_align = recover_unmatched_from_free_onsets(df_align, df_sc, df_on)
    df_align = recover_structured_weak_candidates(df_align, df_sc, df_on, df_weak)
    df_align = recover_platform_weak_bridge(df_align, df_sc, df_weak)
    df_align = anchor_terminal_suffix_to_observed_black(
        df_align, df_sc, df_on, observed_black_time
    )
    df_align = append_duration_ratio_columns(df_align, observed_black_time)
    df_metrics = build_metrics_table(df_align)
    return df_align, df_metrics


def print_change_log():
    print("\n================ 本次 DP3 改动点 ================")
    print("1. 新增 [阶段2] 装饰音后卷修正 ornament_back_refine()")
    print("2. 只处理：前一个与当前同音 + 当前是装饰音 + 后一个是主音")
    print("3. 只允许在 当前红点 ~ 后一个主音红点 之间后卷")
    print("4. 音高规则：<=25c 不罚，25~50c 小惩罚，>50c 直接禁用")
    print("5. 时值评分：前音时值更好 + 装饰音更接近 6/BPM，且装饰音权重更高")
    print("6. 只有新方案整体更好才替换，否则不动")
    print("7. 主流程改为：DP2恢复 -> 装饰音后卷 -> 同音竞争(含尾部black锚点) -> 前卷 -> 输出")
    print("8. 同音竞争区新增主音-主音 tiny-gap 硬约束：同音主音间隔不得小于 50ms 或理论时值25%")
    print("9. 同音竞争评分修复：装饰音->主音若累计时间相同，使用 6/BPM 作为理论间隔，避免 cost 全部为 inf")
    print("10. 尾部改用 onset 表中的真实 black 终点；末小节若有未解决事件，将其后缀按顺序联合锚定，装饰音允许 unmatched")
    print("================================================\n")


def main():
    base_dir = os.path.abspath(
        os.environ.get("ERHU_WORK_DIR", os.path.dirname(os.path.abspath(__file__)))
    )

    print("==== Score-driven Alignment (DP3 simplified version) ====")
    print("method: load DP2 result -> ornament back refine -> same-note competition -> forward early refine -> output")
    print("note: DP3 = 先装饰音后卷，再同音竞争，最后前卷")
    print("============================================================")

    print_change_log()

    score_file = find_score_file(base_dir)
    prefix = get_prefix_from_score(score_file)
    onset_file = find_onset_file(base_dir, prefix)
    dp2_file = find_dp2_align_file(base_dir, prefix)

    score_path = os.path.join(base_dir, score_file)
    onset_path = os.path.join(base_dir, onset_file)
    dp2_path = os.path.join(base_dir, dp2_file)

    out_align_csv = os.path.join(base_dir, f"{prefix}_{RUN_TAG}_对齐结果_{METHOD_TAG_CN}.csv")
    out_metrics_csv = os.path.join(base_dir, f"{prefix}_{RUN_TAG}_对齐指标_{METHOD_TAG_CN}.csv")
    out_excel_xlsx = os.path.join(base_dir, f"{prefix}_{RUN_TAG}_对齐结果_{METHOD_TAG_CN}.xlsx")

    print(f"检测到 SCORE : {score_file}")
    print(f"推断曲名前缀: {prefix}")
    print(f"检测到 ONSET : {onset_file}")
    print(f"检测到 DP2   : {dp2_file}")
    print()

    df_sc, df_on, df_weak, observed_black_time = load_data(score_path, onset_path)
    sc_f = df_sc[SCORE_F0_COL].to_numpy(dtype=float)
    sc_m = midi_round(sc_f)

    print(f"score 音符数 : {len(df_sc)}")
    print(f"onset 候选数 : {len(df_on)}")
    print(f"weak 平台候选数（仅 DP3 局部补救）: {len(df_weak)}")
    print(
        "真实 black 终点: "
        + (f"{observed_black_time:.6f}s" if np.isfinite(observed_black_time) else "未检测到")
    )

    print("读取 DP2 结果并恢复匹配对...")
    df_dp2_for_dp3 = ensure_dp2_columns(pd.read_csv(dp2_path))
    df_dp2_for_dp3, n_single = repair_single_pitch_f(df_dp2_for_dp3, df_on)
    df_dp2_for_dp3 = recalc_dp2_metrics(df_dp2_for_dp3)
    df_dp2_for_dp3, n_triplet = repair_triplet_from_duration_f(df_dp2_for_dp3, df_on)
    df_dp2_for_dp3 = enforce_unique_onset_assignments(df_dp2_for_dp3)
    print(f"[DP3前置精修] 单点F={n_single}, 节奏F三音={n_triplet}")
    pairs = load_pairs_from_dp2(df_dp2_for_dp3, df_sc, df_on)
    print(f"[阶段1] 从 DP2 恢复得到 {len(pairs)} 对")

    pairs = run_dp3(
        pairs,
        df_sc,
        df_on,
        sc_m,
        sc_f,
        observed_black_time=observed_black_time,
    )

    if len(pairs) != len(df_sc):
        print(f"⚠️ 最终匹配数 {len(pairs)}，score 数 {len(df_sc)}")
    else:
        print("✅ 最终匹配数与 score 数一致")

    df_align, df_metrics = run_output_stage(
        df_sc,
        df_on,
        pairs,
        df_weak,
        observed_black_time=observed_black_time,
    )

    df_align.to_csv(out_align_csv, index=False, encoding="utf-8-sig")
    df_metrics.to_csv(out_metrics_csv, index=False, encoding="utf-8-sig")

    with pd.ExcelWriter(out_excel_xlsx, engine="openpyxl") as writer:
        df_align.to_excel(writer, sheet_name="对齐结果", index=False)
        df_metrics.to_excel(writer, sheet_name="对齐指标", index=False)

        ws = writer.book["对齐结果"]
        red_fill = PatternFill(fill_type="solid", fgColor="FFCCCC")
        header = [cell.value for cell in ws[1]]
        status_col_idx = header.index("对齐状态") + 1

        for row_idx in range(2, ws.max_row + 1):
            status_val = ws.cell(row=row_idx, column=status_col_idx).value
            if status_val == "未匹配(后删)":
                for col_idx in range(1, ws.max_column + 1):
                    ws.cell(row=row_idx, column=col_idx).fill = red_fill

    print("输出完成：")
    print("对齐结果表 :", out_align_csv)
    print("对齐指标表 :", out_metrics_csv)
    print("Excel总表  :", out_excel_xlsx)


if __name__ == "__main__":
    main()
